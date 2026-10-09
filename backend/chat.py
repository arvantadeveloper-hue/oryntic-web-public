import json
from datetime import datetime, timezone, timedelta
import re
import logging
import asyncio
import base64
import io
from dataclasses import dataclass, field
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from assistant_persona import persona_block
from db import db, now_iso, new_id, clean
from auth import current_user, workspace_id, _lang_name, lang_rule
from llm import quota_message, llm_text, record_usage, text_credits, describe_image, VISION_CREDITS, quota_exceeded, llm_json
from realtime import notify
from ratelimit import rate_limit, throttle_message, inflight_start, inflight_end
from tools import route_model, wants_tool, plan_tool, GITHUB_RE, SOCIAL_RE, run_image_tool, norm_image_opts, IMAGE_PRESETS, is_media_refusal, EDIT_RE, get_routing, TASK_CONTEXT, REVISE_RE, save_revision, revise_with_llm, plan_task, CAL_RE
from pricing import refresh as pricing_refresh
import seedance
from llm import model_label

MAX_ATTACHMENT_BYTES = 8 * 1024 * 1024
HIST_MSG_CAP = 1500

router = APIRouter(prefix="/api", tags=["chat"])


def _pdf_text(data: str) -> str:
    from pypdf import PdfReader
    raw = base64.b64decode(data.split(",")[-1])
    if len(raw) > MAX_ATTACHMENT_BYTES:
        raise ValueError("attachment too large")
    reader = PdfReader(io.BytesIO(raw))
    return "".join((p.extract_text() or "") + "\n" for p in reader.pages[:20])[:6000]


async def _gallery_task_context(a: dict, user_id: str) -> Optional[str]:
    from files import _same_workspace
    t = await db.tasks.find_one({"id": a["task_id"]}, {"_id": 0, "goal": 1, "final_output": 1, "workspace_id": 1, "user_id": 1})
    if not t or not (t.get("user_id") == user_id or await _same_workspace(user_id, t.get("workspace_id") or "")):
        return None
    return f"[Dokumen Galeri '{t.get('goal') or a.get('name', 'berkas')}' (/workspace/{a['task_id']})]:\n{(t.get('final_output') or '')[:20000]}"


def _blob_text(data: bytes, path: str, name: str) -> str:
    ext = path.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        return f"[PDF '{name}']:\n{_pdf_text(base64.b64encode(data).decode())}"
    if ext == "docx":
        from docx import Document
        return f"[Dokumen '{name}']:\n" + "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)[:6000]
    return f"[Berkas '{name}']:\n{data.decode('utf-8', 'ignore')[:6000]}"


async def _gallery_context(a: dict, user_id: str) -> Optional[str]:
    """Attachment picked from the Gallery: read the stored object / task result — no re-upload."""
    from files import _path_parts, _same_workspace
    from storage import get_object
    if a.get("task_id"):
        return await _gallery_task_context(a, user_id)
    name, path = a.get("name", "berkas"), a.get("path") or ""
    parts = _path_parts(path)
    if parts[2] != user_id and not await _same_workspace(user_id, parts[2]):
        return None
    if a.get("kind") == "video":
        return f"[Video Galeri '{name}' dilampirkan sebagai tautan]"
    data, ctype = await asyncio.to_thread(get_object, path)
    if a.get("kind") != "image" and not (ctype or "").startswith("image/"):
        return _blob_text(data, path, name)
    desc = await describe_image(base64.b64encode(data).decode())
    if not desc:
        return None
    await record_usage(user_id, "vision", VISION_CREDITS, {"name": name})
    return f"[Gambar Galeri '{name}']: {desc}"


async def _drive_context(a: dict, user_id: str) -> Optional[str]:
    """Attachment picked from the user's Google Drive: read it via the Drive API (only works when Drive is connected) — nothing is uploaded to the platform."""
    from integrations import drive_content
    f = await drive_content(user_id, a.get("drive_id") or "")
    name = f.get("name") or a.get("name", "berkas")
    if f.get("image_b64"):
        desc = await describe_image(f["image_b64"])
        if not desc:
            return None
        await record_usage(user_id, "vision", VISION_CREDITS, {"name": name, "drive_id": f.get("id")})
        return f"[Gambar Google Drive '{name}']: {desc}"
    return f"[Berkas Google Drive '{name}' ({f.get('webViewLink', '')})]:\n{(f.get('text') or '')[:6000]}"


async def _attachment_context(a: dict, user_id: str) -> Optional[str]:
    """Text the model should see for one attachment (vision description, PDF text, or raw text)."""
    atype, name, data = a.get("type", "text"), a.get("name", "file"), a.get("data", "")
    if atype == "gallery":
        return await _gallery_context(a, user_id)
    if atype == "drive":
        return await _drive_context(a, user_id)
    if atype == "image":
        desc = await describe_image(data.split(",")[-1])
        if not desc:
            return None
        await record_usage(user_id, "vision", VISION_CREDITS, {"name": name})
        return f"[Gambar '{name}']: {desc}"
    if atype == "pdf":
        return f"[PDF '{name}']:\n{_pdf_text(data)}"
    return f"[Berkas '{name}']:\n{str(data)[:6000]}"


async def _process_attachments(attachments, user_id):
    """Return (context_text, light_meta_list). Extracts text from pdf/text, vision-describes images."""
    ctx, meta = [], []
    for a in (attachments or [])[:5]:
        meta.append({"type": a.get("type", "text"), "name": a.get("name", "file"), **{k: a[k] for k in ("path", "task_id", "kind", "drive_id", "link", "mime") if a.get(k)}})
        try:
            text = await _attachment_context(a, user_id)
        except Exception:
            text = f"[Lampiran '{a.get('name', 'file')}' tidak dapat diproses]"
        if text:
            ctx.append(text)
    return ("\n\n".join(ctx), meta)


class ConvIn(BaseModel):
    persona_ids: list = []
    type: str = "private"  # private | group | meeting
    title: Optional[str] = None
    participant_ids: list = []  # other workspace humans invited to a meeting


def _can_access(conv: dict, u: dict) -> bool:
    if not conv:
        return False
    if conv.get("user_id") == u["id"]:
        return True
    if u["id"] in (conv.get("participants") or []):
        return True
    if u.get("role") == "admin" and conv.get("workspace_id") == workspace_id(u):
        return True
    return False


class MsgIn(BaseModel):
    content: str = Field(min_length=1, max_length=20000)
    attachments: list = []
    moderator: bool = True  # legacy flag (ignored): Moderator now only interjects when the discussion is stuck
    voice_mode: bool = False  # spoken conversation: short, warm, human-like replies (no markdown)
    interrupted: bool = False  # the user barged in while the assistant was speaking
    channel: Optional[str] = Field(default=None, pattern="^(meeting_chat)$")  # typed in the live meeting's chat panel
    reply_to: Optional[dict] = None  # {id, name, content} of the quoted message
    forwarded: bool = False
    model_choice: Optional[str] = Field(default=None, pattern="^[a-z0-9_-]+$")  # user's answer to a model-switch confirmation ("__own__" or a catalog model id)
    replay: bool = False  # re-send of an already stored user message (after choosing the model) → don't store it again
    choice_msg_id: Optional[str] = None  # the model_choice card to remove once answered


class MemIn(BaseModel):
    persona_id: Optional[str] = None
    content: str = Field(min_length=1, max_length=1000)


async def _get_personas(ids, wid=None):
    from support_agent import is_support, support_persona
    out = []
    for pid in ids:
        if is_support(pid):
            sp = await support_persona()
            if sp:
                out.append(sp)
            continue
        q = {"id": pid}
        if wid:
            q["user_id"] = wid
        p = await db.personas.find_one(q, {"_id": 0})
        if p:
            out.append(p)
    return out


# ---------- conversations ----------
def _conv_title(ctype: str, personas: list) -> str:
    names = ", ".join(p["name"] for p in personas)
    if ctype == "meeting":
        return "Panggilan: " + names
    if ctype == "group":
        return "Grup: " + names
    return personas[0]["name"]


async def _resolve_participants(u: dict, ids: list) -> list:
    """Creator + invited friends (deduplicated)."""
    from friends import friend_ids
    participants = [u["id"]]
    if ids:
        allowed = set(await friend_ids(u["id"]))
        participants += [i for i in dict.fromkeys(ids) if i in allowed and i not in participants]
    return participants


def view_title(conv: dict, uid: str) -> dict:
    """Direct human chats are titled with the *other* person's name."""
    t = (conv.get("titles") or {}).get(uid)
    return {**conv, "title": t} if t else conv


@router.post("/conversations")
async def create_conv(x: ConvIn, u: dict = Depends(current_user)):
    from support_agent import SUPPORT_ID
    if SUPPORT_ID in (x.persona_ids or []) and (len(x.persona_ids) > 1 or x.type != "private" or (x.participant_ids or [])):
        raise HTTPException(400, "Oryntix (dukungan) hanya tersedia di percakapan privat")
    personas = await _get_personas(x.persona_ids, workspace_id(u))
    participants = await _resolve_participants(u, x.participant_ids or [])
    if not personas and len(participants) < 2:
        raise HTTPException(400, "Pilih minimal satu asisten atau teman untuk memulai percakapan")
    multi = x.type in ("group", "meeting") and (len(personas) > 1 or x.type == "meeting" or len(participants) > 1)
    ctype = x.type if multi else "private"
    humans = [h async for h in db.users.find({"id": {"$in": participants}}, {"_id": 0, "id": 1, "name": 1, "email": 1, "avatar": 1})] if len(participants) > 1 else []
    doc = {
        "id": new_id(), "user_id": u["id"], "workspace_id": workspace_id(u),
        "participants": participants, "humans": humans,
        "type": ctype,
        "persona_ids": [p["id"] for p in personas],
        "persona_id": personas[0]["id"] if personas else None,
        "members": [{"id": p["id"], "name": p["name"], "portrait": p.get("portrait"), "voice": p.get("voice", "alloy"), **({"builtin": True} if p.get("builtin") else {})} for p in personas],
        "title": x.title or (_conv_title(ctype, personas) if personas else "Grup " + ", ".join(h["name"] for h in humans)), "created_at": now_iso(), "updated_at": now_iso(), "last_message": "",
    }
    await db.conversations.insert_one(dict(doc))
    return clean(doc)


@router.get("/conversations")
async def list_conv(q: Optional[str] = None, limit: int = Query(200, ge=1, le=200), offset: int = Query(0, ge=0), u: dict = Depends(current_user)):
    query = {"$or": [{"user_id": u["id"]}, {"participants": u["id"]}], "archived_conv": {"$ne": True}}
    if q:
        query["title"] = {"$regex": q, "$options": "i"}
    rows = await db.conversations.find(query, {"_id": 0, "invite_token": 0}).sort("updated_at", -1).skip(offset).to_list(limit)
    out = []
    for c in rows:
        c["unread"] = bool(c.get("last_message")) and (c.get("updated_at") or "") > ((c.get("read_at") or {}).get(u["id"]) or "")
        out.append(view_title(c, u["id"]))
    return out


class DirectIn(BaseModel):
    persona_id: str


@router.post("/conversations/direct")
async def direct_conv(x: DirectIn, u: dict = Depends(current_user)):
    """WhatsApp-style: exactly one private chat per assistant (created on first open)."""
    personas = await _get_personas([x.persona_id], workspace_id(u))
    if not personas:
        raise HTTPException(404, "Asisten tidak ditemukan")
    p = personas[0]
    conv = await db.conversations.find_one({"user_id": u["id"], "type": "private", "persona_ids": [p["id"]], "archived_conv": {"$ne": True}}, {"_id": 0}, sort=[("updated_at", -1)])
    if not conv:
        conv = {"id": new_id(), "user_id": u["id"], "workspace_id": workspace_id(u), "participants": [u["id"]], "type": "private", "persona_ids": [p["id"]], "persona_id": p["id"],
                "members": [{"id": p["id"], "name": p["name"], "portrait": p.get("portrait"), "voice": p.get("voice", "alloy")}], "title": p["name"],
                "created_at": now_iso(), "updated_at": now_iso(), "last_message": ""}
        await db.conversations.insert_one(dict(conv))
    return clean(conv)


@router.post("/conversations/{cid}/read")
async def mark_read(cid: str, u: dict = Depends(current_user)):
    await db.conversations.update_one({"id": cid}, {"$set": {f"read_at.{u['id']}": now_iso()}})
    return {"ok": True}


@router.delete("/conversations/{cid}/task")
async def detach_task(cid: str, u: dict = Depends(current_user)):
    """Stop using a Workspace task as this chat's shared context."""
    conv = await db.conversations.find_one({"id": cid})
    if not _can_access(conv, u):
        raise HTTPException(404, "Conversation not found")
    await db.conversations.update_one({"id": cid}, {"$set": {"task_id": None}})
    return {"ok": True}


async def migrate_direct_chats():
    """One private chat per (user, assistant): older duplicates are merged into the newest one as archived history."""
    async for cv in db.conversations.find({"read_at": {"$exists": False}}, {"_id": 0, "id": 1, "user_id": 1, "updated_at": 1}):
        await db.conversations.update_one({"id": cv["id"]}, {"$set": {f"read_at.{cv['user_id']}": cv.get("updated_at") or now_iso()}})
    pipeline = [{"$match": {"type": "private", "archived_conv": {"$ne": True}, "persona_ids": {"$size": 1}}},
                {"$group": {"_id": {"u": "$user_id", "p": {"$arrayElemAt": ["$persona_ids", 0]}}, "n": {"$sum": 1}}}, {"$match": {"n": {"$gt": 1}}}]
    async for g in db.conversations.aggregate(pipeline):
        convs = await db.conversations.find({"user_id": g["_id"]["u"], "type": "private", "persona_ids": [g["_id"]["p"]], "archived_conv": {"$ne": True}}, {"_id": 0}).sort("updated_at", -1).to_list(200)
        primary, old = convs[0], convs[1:]
        for c in old:
            msgs = await db.messages.find({"conversation_id": c["id"]}, {"_id": 0, "id": 1, "created_at": 1, "content": 1}).sort("created_at", 1).to_list(5000)
            if msgs:
                await db.messages.update_many({"conversation_id": c["id"]}, {"$set": {"conversation_id": primary["id"], "archived": True, "origin_conversation_id": c["id"]}})
                await db.chat_archives.insert_one({"id": new_id(), "user_id": c["user_id"], "conversation_id": primary["id"], "title": c.get("title") or primary["title"],
                                                   "persona_names": [m["name"] for m in c.get("members") or []], "period_start": msgs[0]["created_at"], "period_end": msgs[-1]["created_at"],
                                                   "summary": c.get("memory_summary") or (msgs[-1].get("content") or "")[:400], "message_ids": [m["id"] for m in msgs], "message_count": len(msgs),
                                                   "reason": "merged", "restored": False, "created_at": now_iso()})
            await db.conversations.update_one({"id": c["id"]}, {"$set": {"archived_conv": True, "merged_into": primary["id"]}})
        if primary.get("title", "").startswith("Chat dengan "):
            await db.conversations.update_one({"id": primary["id"]}, {"$set": {"title": primary["title"][len("Chat dengan "):]}})


@router.get("/conversations/{cid}/messages")
async def get_messages(cid: str, before: Optional[str] = None, limit: int = Query(50, ge=1, le=200), archived: int = 0, u: dict = Depends(current_user)):
    """Newest `limit` messages (ascending). `before` = created_at cursor for scrolling up; archived=1 shows the archive."""
    conv = await db.conversations.find_one({"id": cid}, {"_id": 0, "invite_token": 0})
    if not _can_access(conv, u):
        raise HTTPException(404, "Conversation not found")
    query = {"conversation_id": cid, "archived": True} if archived else {"conversation_id": cid, "archived": {"$ne": True}}
    if before:
        query["created_at"] = {"$lt": before}
    page = await db.messages.find(query, {"_id": 0}).sort("created_at", -1).to_list(limit + 1)
    has_more = len(page) > limit
    msgs = list(reversed(page[:limit]))
    out = {"conversation": view_title(conv, u["id"]), "messages": msgs, "has_more": has_more}
    if not before and not archived:
        out["archived_count"] = await db.messages.count_documents({"conversation_id": cid, "archived": True})
        out["archives_count"] = await db.chat_archives.count_documents({"conversation_id": cid})
        out["long_chat"] = await _long_chat(cid, conv)
    return out


LONG_CHAT_MSGS, LONG_CHAT_CHARS = 40, 15000


async def _long_chat(cid: str, conv: dict) -> bool:
    """True when the live (non-archived) thread is long enough that the assistant should offer to summarize."""
    live = await db.messages.find({"conversation_id": cid, "archived": {"$ne": True}}, {"_id": 0, "content": 1}).to_list(500)
    snooze = int((conv or {}).get("summary_snoozed_at_count") or 0)
    return len(live) - snooze >= LONG_CHAT_MSGS or sum(len(m.get("content") or "") for m in live[snooze:]) >= LONG_CHAT_CHARS


CORE_SUMMARY_CAP = 1500
PERIOD_SUMMARY_CAP = 1200


async def _merge_summary(cid: str, u: dict, conv: dict, instruction: str) -> tuple:
    """Two tiers: (core, detailed, credits). Core = always-on facts/decisions/preferences (<=1500 chars, merged with the
    previous core); detailed = full period summary stored with the archive and only injected when relevant."""
    history = await _history_text(cid, limit=120, with_summary=False)
    prev = conv.get("memory_summary") or ""
    lang = _lang_name(u)
    detailed = await llm_text(f"You write entirely in {lang}. {instruction} Output markdown, compact but complete; keep every decision, number, name, deadline and open question.",
                              f"Discussion:\n{history}\n\nDetailed summary:")
    core_prompt = (f"Earlier core memory:\n{prev}\n\n" if prev else "") + f"New period summary:\n{detailed}\n\nUpdated core memory:"
    core = await llm_text(f"You write entirely in {lang}. Produce the CORE MEMORY of this conversation: only durable facts, decisions, user preferences, names, numbers and open commitments, "
                          f"as terse bullet points, max {CORE_SUMMARY_CAP} characters. MERGE the earlier core memory with the new period summary; drop chit-chat and superseded details.",
                          core_prompt)
    core = core[:CORE_SUMMARY_CAP]
    return core, detailed, text_credits(history, detailed) + text_credits(core_prompt, core)


def _keywords(text: str) -> set:
    return {w.strip(".,?!:;\"'()").lower() for w in (text or "").split() if len(w.strip(".,?!:;\"'()")) > 3}


async def _relevant_periods(cid: str, query: str, cap: int = 2) -> list:
    """Detailed period summaries from this chat's archives that share keywords with the current message."""
    kws = _keywords(query)
    if not kws:
        return []
    out = []
    async for a in db.chat_archives.find({"conversation_id": cid, "restored": {"$ne": True}}, {"_id": 0, "summary": 1, "period_start": 1, "period_end": 1}).sort("created_at", -1).limit(12):
        score = len(kws & _keywords(a.get("summary") or ""))
        if score >= 2:
            out.append((score, a))
    out.sort(key=lambda t: -t[0])
    return [a for _, a in out[:cap]]


@router.post("/conversations/{cid}/compact")
async def compact_conversation(cid: str, u: dict = Depends(current_user)):
    """Summarize the live thread into conversation memory and archive the raw messages (kept in Arsip, restorable)."""
    from archives import archive_conversation
    conv = await db.conversations.find_one({"id": cid}, {"_id": 0})
    if not _can_access(conv, u):
        raise HTTPException(404, "Conversation not found")
    entry = await archive_conversation(conv, u, "manual")
    if not entry:
        raise HTTPException(400, "Tidak ada pesan untuk dirangkum")
    return {"summary": entry["summary"], "archived": entry["message_count"], "credits_used": 0, "archive_id": entry["id"]}


@router.post("/conversations/{cid}/summary-later")
async def summary_later(cid: str, u: dict = Depends(current_user)):
    conv = await db.conversations.find_one({"id": cid}, {"_id": 0})
    if not _can_access(conv, u):
        raise HTTPException(404, "Conversation not found")
    live = await db.messages.count_documents({"conversation_id": cid, "archived": {"$ne": True}})
    await db.conversations.update_one({"id": cid}, {"$set": {"summary_snoozed_at_count": live}})
    return {"ok": True}


DEFAULT_NOTULEN_FIELDS = [{"name": "Agenda", "required": True}, {"name": "Pembahasan", "required": True}, {"name": "Keputusan", "required": True},
                          {"name": "Tindak lanjut (PIC & tenggat)", "required": True}, {"name": "Isu terbuka", "required": False}]


async def _notulen_fields(u: dict) -> list:
    owner = await db.users.find_one({"id": workspace_id(u)}, {"_id": 0, "settings": 1}) or {}
    fields = (owner.get("settings") or {}).get("notulen_fields")
    return fields if fields else DEFAULT_NOTULEN_FIELDS


@router.post("/conversations/{cid}/notulen-check")
async def notulen_check(cid: str, u: dict = Depends(current_user)):
    """Which notulen fields are still empty? Used before 'Akhiri & Simpan Notulen'."""
    conv = await db.conversations.find_one({"id": cid}, {"_id": 0})
    if not _can_access(conv, u):
        raise HTTPException(404, "Conversation not found")
    fields = await _notulen_fields(u)
    history = await _history_text(cid, limit=60)
    if not history.strip():
        return {"fields": fields, "missing": [f["name"] for f in fields if f.get("required")], "notes": {}}
    sys = ('You audit meeting minutes. For each field decide if the discussion contains enough concrete content to fill it. '
           'Reply JSON only: {"fields": {"<field name>": {"filled": bool, "note": "<what is missing, in ' + _lang_name(u) + ', max 15 words>"}}}')
    try:
        res = await llm_json(sys, f"Fields: {json.dumps([f['name'] for f in fields], ensure_ascii=False)}\n\nDiscussion:\n{history}")
    except Exception:
        res = {}
    status = res.get("fields") or {}
    missing = [f["name"] for f in fields if f.get("required") and not (status.get(f["name"]) or {}).get("filled")]
    return {"fields": fields, "missing": missing, "notes": {k: v.get("note", "") for k, v in status.items() if isinstance(v, dict)}}


@router.delete("/conversations/{cid}")
async def del_conv(cid: str, u: dict = Depends(current_user)):
    conv = await db.conversations.find_one({"id": cid})
    if not conv:
        return {"ok": True}
    if conv.get("user_id") != u["id"] and u.get("role") != "admin":
        # a non-owner participant just leaves the conversation
        await db.conversations.update_one({"id": cid}, {"$pull": {"participants": u["id"]}})
        return {"ok": True, "left": True}
    await db.conversations.delete_one({"id": cid})
    await db.messages.delete_many({"conversation_id": cid})
    return {"ok": True}


# Shared core character of EVERY assistant (placeholders filled per persona/user). Set by the platform owner.
NO_REPEAT_TEXT = ("Read what the other assistants already answered in this thread. Do NOT repeat or rephrase their points; "
                  "agree in a few words if needed and add something NEW, or say briefly you have nothing to add.")

MEETING_CHAT_STYLE = ("MEETING CHAT PANEL: a live voice meeting is in progress and the user just TYPED this message in the meeting's text "
                      "chat panel. Reply in TEXT only (this reply is shown in the panel, not spoken): use markdown freely — tables, "
                      "numbered lists, code blocks, links — whenever it makes the data clearer. Be complete but compact; do not greet, "
                      "do not say you will 'speak' or 'read' anything aloud.")


def _relevant_memories(mems: list, query: str, cap: int = 10) -> list:
    """Always-on (pinned) memories + those sharing words with the current message; everything if the list is small."""
    if len(mems) <= cap:
        return mems
    if not query:
        return [m for m in mems if m.get("pinned")] + [m for m in mems if not m.get("pinned")][:cap]
    words = _keywords(query)
    pinned = [m for m in mems if m.get("pinned")]
    rest = sorted([m for m in mems if not m.get("pinned")], key=lambda m: -len(words & _keywords(m["content"])))
    return pinned + rest[:max(0, cap - len(pinned))]


async def _persona_system(persona, user, roster=None, voice_mode=False, query=None):
    prof = persona.get("profile", {})
    lang_name = _lang_name(user)
    parts = [f"CRITICAL: You MUST always write every reply in {lang_name}, no matter what language these instructions or the persona profile are written in, and regardless of which AI model is answering or which tool/task mode is active. Never switch to another language unless the user explicitly asks for it."]
    parts.append(f"You are '{persona['name']}', an AI Assistent. {prof.get('system_instructions','')}")
    parts += persona_block(persona["name"], lang_name, voice_mode)
    if persona.get("builtin"):
        from support_agent import support_prompt_parts
        parts += await support_prompt_parts()
        parts.append(lang_rule(user))
        return "\n".join(parts)
    parts.append(f"VIDEO PRICING (platform credits, from the admin price list): {await video_pricing_text()}; default clip {seedance.DEFAULT_DUR} s at 720p, Seedance 2.0 up to 15 s, Seedance 2.5 up to 30 s; resolution 480p (×{seedance.multipliers()['res'].get('480p')}) / 720p / 1080p (×{seedance.multipliers()['res'].get('1080p')}); real-person mode (Seedance 2.0, image-to-video only, ×{seedance.multipliers()['real_person']}); generated sound/ambience ×{seedance.multipliers()['audio']}. "
                 "Rendered videos are saved to the user's Google Drive (must be connected). When the user asks how much a video / Seedance costs, quote exactly these credits per second and the total for their duration.")
    parts.append("CAPABILITIES: You CAN create and show images (photorealistic photos, renders, illustrations, logos, posters), short videos/clips and downloadable documents directly in this chat — the platform renders them for you automatically whenever the user asks. NEVER say you cannot render, generate, display or send images or videos. If the user asks for one and it has not appeared yet, simply say briefly that you are preparing it.")
    mems = _relevant_memories(await db.memory_items.find({"user_id": user["id"], "persona_id": persona["id"], "enabled": True}).to_list(50), query)
    if mems:
        parts.append("Saved memory about the user: " + "; ".join(m["content"] for m in mems))
    from provider_tools import uses_file_search
    if query and not uses_file_search(persona, persona.get("model")):  # with file_search the model retrieves from the vector store itself
        from knowledge import relevant_knowledge
        kn = await relevant_knowledge(persona["id"], query)
        if kn:
            parts.append("Reference knowledge relevant to this message (from the assistant's uploaded documents; cite the title when you use it):\n" + "\n".join(f"[{x['title']}] {x['text']}" for x in kn))
    if roster:
        others = [n for n in roster if n != persona["name"]]
        if others:
            parts.append(f"You are in a group conversation with the user and other AI assistants: {', '.join(others)}. "
                         f"Respond only as {persona['name']}, keep it concise, build on what others said without repeating them, and do not speak for the others.")
    parts.append(lang_rule(user))
    return "\n".join(parts)


async def _history_text(cid: str, limit=14, with_summary=True, query: str = "") -> str:
    msgs = await db.messages.find({"conversation_id": cid, "archived": {"$ne": True}, "is_summary": {"$ne": True}}, {"_id": 0}).sort("created_at", 1).to_list(1000)
    lines = []
    if with_summary:
        conv = await db.conversations.find_one({"id": cid}, {"_id": 0, "memory_summary": 1}) or {}
        if conv.get("memory_summary"):
            lines.append(f"[Memori inti percakapan]: {conv['memory_summary'][:CORE_SUMMARY_CAP + 500]}")
        for a in await _relevant_periods(cid, query):
            lines.append(f"[Rincian periode {a.get('period_start', '')[:10]}–{a.get('period_end', '')[:10]}, relevan dengan pertanyaan]: {(a.get('summary') or '')[:PERIOD_SUMMARY_CAP]}")
    for m in msgs[-limit:]:
        if m["role"] == "user":
            who = m.get("sender_name") or "User"
        else:
            who = m.get("persona_name") or "Assistant"
        body = m["content"] or ""
        if m["role"] != "user" and len(body) > HIST_MSG_CAP:
            body = body[:HIST_MSG_CAP] + " …(dipangkas; versi lengkap tersimpan di Ruang Kerja/Galeri)"
        lines.append(f"{who}: {body}")
        if m.get("attachment_text"):
            lines.append(f"[Isi lampiran {who}]: {m['attachment_text'][:1500]}")
        if m.get("media"):
            lines.append("[Berkas yang dibuat asisten]: " + ", ".join(x.get("name", "") for x in m["media"]))
    return "\n".join(lines)


async def _load_ai_conv(cid: str, u: dict):
    """Access + quota + persona checks shared by the AI endpoints. Returns (conv, personas)."""
    conv = await db.conversations.find_one({"id": cid})
    if not _can_access(conv, u):
        raise HTTPException(404, "Conversation not found")
    over = await quota_exceeded(u)
    if over:
        raise HTTPException(402, quota_message(over))
    personas = await _get_personas(conv.get("persona_ids", []))
    if not personas and len(conv.get("participants") or []) < 2:
        raise HTTPException(400, "Percakapan ini tidak memiliki persona")
    return conv, personas


def _mentioned(content: str, personas: list) -> list:
    """@mention routing: personas explicitly named with @ in the message."""
    lower = content.lower()
    return [p for p in personas if ("@" + p["name"].lower().replace(" ", "")) in lower.replace(" ", "") or ("@" + p["name"].lower()) in lower]


async def _owner_settings(u: dict) -> dict:
    owner = await db.users.find_one({"id": workspace_id(u)}, {"_id": 0, "settings": 1}) or {}
    return owner.get("settings") or {}


async def _save_ai_msg(cid: str, persona: dict, content: str, used: int, via, extra: dict) -> dict:
    ai_msg = {"id": new_id(), "conversation_id": cid, "role": "assistant", "content": content,
              "persona_id": persona["id"], "persona_name": persona["name"],
              "portrait": persona.get("portrait"), "credits": used, "created_at": now_iso(), **extra}
    if via:
        ai_msg["via"] = via
    await db.messages.insert_one(dict(ai_msg))
    if via != "realtime" and not (extra or {}).get("tool"):  # tool/system cards (task done, search results) already notify on their own channel
        await _fanout_message(cid, "__assistant__", persona["name"], content)
    return ai_msg


async def _emit_final(ctx, text: str, credits: int, extra: dict):
    """Persist the assistant message for a tool turn and yield its final SSE + credits."""
    msg = await _save_ai_msg(ctx.cid, ctx.persona, text, credits, ctx.via, extra)
    payload = {"message_id": msg["id"], "content": text}
    for k in ("media", "pending_tool", "task_id", "task_version", "tool", "pending_task", "results", "cta"):
        if k in extra:
            payload[k] = extra[k]
    yield ctx.sse(final=True, **payload)
    await notify(ctx.cid, {"type": "message", "role": "assistant", "persona_id": ctx.persona["id"], "message": clean(msg)})
    yield credits


async def _run_image_set(task_id: str, uid: str, prompts: list, cid: str, persona: dict):
    """Background: generate each image of a multi-image request into a Workspace task (progress visible in the task panel)."""
    from realtime import notify_user
    media, total = [], 0
    for i, prompt in enumerate(prompts):
        await db.tasks.update_one({"id": task_id}, {"$set": {f"steps.{i}.status": "running", "updated_at": now_iso()}})
        try:
            out = await run_image_tool(uid, prompt)
            out["media"][0]["name"] = f"gambar-{i + 1}.{out['media'][0]['name'].rsplit('.', 1)[-1]}"
            media += out["media"]
            total += out["credits"]
            await record_usage(uid, "image_generation", out["credits"], {"conversation_id": cid, "persona_id": persona["id"], "task_id": task_id})
            upd = {f"steps.{i}.status": "completed", f"steps.{i}.output": f"Gambar {i + 1} selesai ({out['credits']} kredit).", "media": media, "credits_used": total, "updated_at": now_iso()}
        except Exception:
            upd = {f"steps.{i}.status": "failed", f"steps.{i}.output": "Gambar ini gagal dibuat.", "updated_at": now_iso()}
        await db.tasks.update_one({"id": task_id}, {"$set": upd})
        await notify_user(uid, {"type": "task_update", "task_id": task_id, "status": "running"})
    lines = "\n".join(f"{i + 1}. {p}" for i, p in enumerate(prompts))
    final = f"## {len(media)} dari {len(prompts)} gambar selesai\n\nPrompt yang dipakai:\n{lines}"
    await db.tasks.update_one({"id": task_id}, {"$set": {"status": "completed" if media else "failed", "final_output": final, "error": None if media else "Semua gambar gagal dibuat", "updated_at": now_iso()}})
    await notify_user(uid, {"type": "task_update", "task_id": task_id, "status": "completed" if media else "failed"})


async def _start_image_set(ctx_user: dict, persona: dict, cid: str, prompts: list, title: str) -> dict:
    task = {"id": new_id(), "user_id": ctx_user["id"], "workspace_id": workspace_id(ctx_user), "goal": title, "type": "image_set", "status": "running",
            "steps": [{"id": new_id(), "role": "Visual", "title": f"Gambar {i + 1}: {p[:90]}", "status": "queued"} for i, p in enumerate(prompts)],
            "summary": f"{len(prompts)} gambar diminta dari chat", "model": "gemini-image", "final_output": "", "media": [], "credits_used": 0,
            "persona_id": persona["id"], "persona_name": persona["name"], "source": "chat", "conversation_ids": [cid], "image_prompts": prompts,
            "version": 1, "created_at": now_iso(), "updated_at": now_iso()}
    await db.tasks.insert_one(dict(task))
    asyncio.create_task(_run_image_set(task["id"], ctx_user["id"], prompts, cid, persona))
    return task


def _image_set_text(task: dict, n: int) -> str:
    return (f"Oke, {n} gambar sekaligus — aku kerjakan sebagai tugas **{task['goal']}** di Ruang Kerja supaya progresnya bisa kamu pantau. "
            f"Klik kartu tugas di bawah untuk melihat hasilnya satu per satu — [buka di Ruang Kerja](/workspace/{task['id']}).")


ANIMATE_RE = re.compile(r"\b(animasikan|animasi|gerakkan|hidupkan|jadikan video|bikin video|buat video|animate|turn .* into (a )?video)\b", re.I)
from chat_media import (_image_turn, _video_turn, _document_turn, _drive_turn, video_offer, video_pricing_text, _run_video_bg, _owner_balance,  # noqa: E402,F401
                        VIDEO_WAIT_TEXT, VIDEO_DONE_TEXT, VIDEO_FAIL_TEXT, IMAGE_DONE_TEXT, IMAGE_EDIT_DONE_TEXT)


def _repo_provider(name: str) -> dict:
    """Git hosting providers share one chat flow; only the API module and wording differ."""
    if name == "gitlab":
        import gitlab as m
        return {"label": "GitLab", "pr": "Merge request", "connected": m.gl_connected, "repos": m.gl_projects, "tree": m.gl_tree, "read": m.gl_read, "issues": m.gl_issues, "create": m.gl_create_mr, "commit": m.gl_commit, "diff": m.gl_mr_diff}
    import github as m
    return {"label": "GitHub", "pr": "Pull request", "connected": m.gh_connected, "repos": m.gh_repos, "tree": m.gh_tree, "read": m.gh_read, "issues": m.gh_issues, "create": m.gh_create_pr, "commit": m.gh_commit, "diff": m.gh_pr_diff}


REVIEW_SYS = ("You are a meticulous senior code reviewer. Review the pull/merge request below and answer in the user's language ({lang}), in markdown, spoken-friendly but precise:\n"
              "1. **Ringkasan** — what the change does (2-4 sentences).\n2. **Risiko & bug potensial** — concrete issues with file + line hints from the diff (security, logic, error handling, performance, breaking changes).\n"
              "3. **Saran perbaikan** — actionable, with short code snippets when useful.\n4. **Kualitas** — naming, tests, docs, style (brief).\n5. **Verdict** — one of: Siap merge / Merge dengan catatan / Perlu perbaikan, with one-line reason.\n"
              "Only comment on what is in the diff; do not invent files. Be specific, skip generic advice.")


def _pr_number(v) -> Optional[int]:
    v = str(v or "").strip().lstrip("#!")
    return int(v) if v.isdigit() else None


async def git_review_text(prov: dict, d: dict, u: dict, model_key: str, request: str = "") -> tuple:
    """LLM code review of a PR/MR diff → (markdown, credits). Shared by chat tool and voice endpoint."""
    diff_text = "\n\n".join(f"--- {f['path']} ({f['status']}) ---\n{f['patch']}" for f in d["files"])
    review = await llm_text(REVIEW_SYS.format(lang=_lang_name(u)), f"{prov['pr']} !{d['number']} — {d['title']} (by {d['author']}, {d['head']} → {d['base']})\nDescription:\n{d['body']}\n\nUser's request: {request or 'review this'}\n\nDIFF:\n{diff_text}", model_key)
    text = f"### Review {prov['pr'].lower()} [!{d['number']} {d['title']}]({d['url']})\n`{d['head']}` → `{d['base']}` · {d['changed_files']} file{' (diff dipotong)' if d['truncated'] else ''}\n\n{review}"
    return text, text_credits(diff_text, review, model_key)


class GitReviewIn(BaseModel):
    provider: str = Field(default="github", max_length=10)
    repo: str = Field(max_length=200)
    number: Optional[int] = None


@router.post("/conversations/{cid}/git-review")
async def conv_git_review(cid: str, x: GitReviewIn, u: dict = Depends(current_user)):
    """Voice tool: review a PR/MR and drop the written review into the conversation as an assistant message."""
    conv = await db.conversations.find_one({"id": cid}, {"_id": 0})
    if not conv or not _can_access(conv, u):
        raise HTTPException(404, "Conversation not found")
    prov = _repo_provider(x.provider)
    if not await prov["connected"](u["id"]):
        raise HTTPException(400, f"{prov['label']} belum terhubung.")
    d = await prov["diff"](u["id"], x.repo, x.number)
    persona = (await _get_personas([conv.get("persona_id") or (conv.get("persona_ids") or [None])[0]]) or [None])[0]
    text, credits = await git_review_text(prov, d, u, (persona or {}).get("model"))
    if persona:
        await _save_ai_msg(cid, persona, text, credits, "text", {"tool": f"{x.provider}_review", "review_of": {"repo": x.repo, "number": d["number"], "url": d["url"]}})
        if credits:
            await record_usage(u["id"], "chat", credits, {"conversation_id": cid, "persona_id": persona["id"], "tool": f"{x.provider}_review"})
    return {"number": d["number"], "title": d["title"], "url": d["url"], "summary": text[:1200]}

KIND_ID = {"text": "teks", "image": "gambar", "video": "video"}
SOCIAL_LABEL = {"linkedin": "LinkedIn", "meta": "Facebook/Instagram", "youtube": "YouTube"}


async def _latest_media(cid: str, kind: str) -> Optional[dict]:
    want = "video" if kind == "video" else "image"
    async for m in db.messages.find({"conversation_id": cid, "media.0": {"$exists": True}}, {"_id": 0, "media": 1}).sort("created_at", -1).limit(30):
        for x in reversed(m.get("media") or []):
            if (x.get("type") or "").startswith(want) and (x.get("path") or (want == "video" and x.get("drive_id"))):
                return x
    return None


async def social_publish_from_chat(cid: str, u: dict, providers: list, caption: str, kind: str, app_url: str, title: Optional[str] = None) -> dict:
    """Resolve the latest image/video of the conversation and publish it (chat run-tool + voice tool)."""
    from social import publish, PublishIn
    media = await _latest_media(cid, kind) if kind in ("image", "video") else None
    if kind in ("image", "video") and not media:
        raise HTTPException(400, f"Belum ada {kind} di percakapan ini untuk diposting.")
    res = await publish(u["id"], PublishIn(providers=providers, kind=kind, text=caption, title=title or (media or {}).get("name") or "", media_path=(media or {}).get("path"), drive_id=(media or {}).get("drive_id"), app_url=app_url, source={"conversation_id": cid}))
    lines = [f"- **{SOCIAL_LABEL.get(r['provider'], r['provider'])}**: " + (f"terkirim — [lihat post]({r.get('post_url')})" + (f" · [Instagram]({r['instagram_url']})" if r.get("instagram_url") else "") if r["status"] == "sent" else f"gagal — {r.get('error')}") for r in res]
    return {"results": res, "text": "Hasil publikasi:\n" + "\n".join(lines) + "\n\nSemua riwayat ada di menu [Social Media](/social)."}


async def _social_turn(ctx, plan: dict):
    from social import accounts
    yield ctx.sse(start=True)
    accs = {a["id"]: a for a in await accounts(ctx.user["id"])}
    providers = [p for p in (plan.get("providers") or []) if p in SOCIAL_LABEL] or [p for p, a in accs.items() if a["connected"]][:1]
    missing = [SOCIAL_LABEL[p] for p in providers if not accs.get(p, {}).get("connected")]
    if not providers or missing:
        txt = (f"Akun {', '.join(missing)} belum terhubung." if missing else "Belum ada akun media sosial yang terhubung.") + " Hubungkan dulu di menu [Social Media](/social), lalu minta lagi ya."
        async for ev in _emit_final(ctx, txt, 0, {"tool": "social_publish", "error": True}):
            yield ev
        return
    kind = plan.get("content_kind") or "image"
    media = await _latest_media(ctx.cid, kind) if kind in ("image", "video") else None
    if kind in ("image", "video") and not media:
        async for ev in _emit_final(ctx, f"Aku belum menemukan {kind} di percakapan ini. Buat dulu atau lampirkan, lalu minta posting lagi.", 0, {"tool": "social_publish", "error": True}):
            yield ev
        return
    caption = (plan.get("caption") or "").strip() or (media or {}).get("name") or ""
    when = _schedule_utc(plan.get("schedule_at"), ctx.user)
    if plan.get("schedule_at") and not when:
        async for ev in _emit_final(ctx, "Waktunya sudah lewat atau tidak kubaca dengan jelas — sebutkan tanggal & jamnya ya (mis. \"besok jam 09.00\").", 0, {"tool": "social_publish", "error": True}):
            yield ev
        return
    pt = {"kind": "social", "providers": providers, "caption": caption, "content_kind": kind, "credits": 0, "media_path": (media or {}).get("path"), "media_name": (media or {}).get("name"), "media_type": (media or {}).get("type")}
    if when:
        pt["schedule_at"], pt["schedule_label"] = when
        text = f"Siap **menjadwalkan** posting {KIND_ID.get(kind, kind)} ke **{', '.join(SOCIAL_LABEL[p] for p in providers)}** pada **{when[1]}** dengan caption:\n\n> {caption}\n\nLanjutkan?"
    else:
        text = f"Siap posting {KIND_ID.get(kind, kind)} ke **{', '.join(SOCIAL_LABEL[p] for p in providers)}** dengan caption:\n\n> {caption}\n\nLanjutkan?"
    yield ctx.sse(delta=text)
    async for ev in _emit_final(ctx, text, 0, {"pending_tool": pt}):
        yield ev


def _schedule_utc(local: str | None, user: dict):
    """'YYYY-MM-DDTHH:MM' in the user's timezone → (UTC ISO, human label) or None when empty/invalid/past."""
    if not local:
        return None
    from zoneinfo import ZoneInfo
    tz = (user.get("settings") or {}).get("timezone") or "Asia/Jakarta"
    try:
        z = ZoneInfo(tz)
    except Exception:
        z = ZoneInfo("Asia/Jakarta")
    try:
        dt = datetime.fromisoformat(local[:16]).replace(tzinfo=z)
    except Exception:
        return None
    if dt <= datetime.now(timezone.utc) + timedelta(minutes=1):
        return None
    hari = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"][dt.weekday()]
    bulan = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"][dt.month - 1]
    return dt.astimezone(timezone.utc).isoformat(), f"{hari}, {dt.day} {bulan} {dt.year} {dt:%H:%M} ({tz})"


async def social_schedule_from_chat(cid: str, u: dict, pt: dict, app_url: str, persona_id, persona_name, portrait) -> dict:
    from social import publish as _p, PublishIn, schedule_post  # noqa: F401
    kind = pt.get("content_kind") or "image"
    media = await _latest_media(cid, kind) if kind in ("image", "video") else None
    if kind in ("image", "video") and not media:
        raise HTTPException(400, f"Belum ada {kind} di percakapan ini untuk dijadwalkan.")
    x = PublishIn(providers=pt["providers"], kind=kind, text=pt.get("caption") or "", title=pt.get("title") or (media or {}).get("name") or "", media_path=(media or {}).get("path"),
                  drive_id=(media or {}).get("drive_id"), app_url=app_url or "https://oryntix.app", source={"conversation_id": cid, "scheduled": True})
    job = await schedule_post(u["id"], x, pt["schedule_at"], pt.get("schedule_label") or pt["schedule_at"],
                              {"conversation_id": cid, "persona_id": persona_id, "persona_name": persona_name, "portrait": portrait})
    return {"job": job, "text": f"Dijadwalkan ✅ — {KIND_ID.get(kind, kind)} akan diposting ke **{', '.join(SOCIAL_LABEL.get(p, p) for p in pt['providers'])}** pada **{job['scheduled_label']}**.\n\n"
                                f"> {pt.get('caption') or ''}\n\nAku akan mengabari di chat ini setelah terkirim. Lihat atau batalkan jadwal di menu [Social Media](/social)."}


class SocialChatIn(BaseModel):
    providers: list[str] = Field(min_length=1, max_length=3)
    caption: str = Field(default="", max_length=3000)
    kind: str = Field(default="image", pattern="^(text|image|video)$")
    app_url: str = Field(min_length=8, max_length=300)


class VoiceImageIn(BaseModel):
    prompt: str = Field(min_length=1, max_length=2000)
    edit_previous: bool = False
    aspect_ratio: str = Field(default="1:1", max_length=8)
    quality: str = Field(default="standar", max_length=12)
    request: str = Field(default="", max_length=300)


class VoiceVideoIn(BaseModel):
    prompt: str = Field(min_length=1, max_length=2000)
    duration: int = Field(default=5, ge=3, le=30)
    aspect_ratio: str = Field(default="16:9", max_length=8)
    resolution: str = Field(default="720p", max_length=8)
    real_person: bool = False
    with_audio: bool = False
    from_image: bool = False
    request: str = Field(default="", max_length=300)


async def _voice_persona(cid: str, u: dict) -> dict:
    conv, personas = await _load_ai_conv(cid, u)
    return personas[0] if personas else {"id": "__system__", "name": "Asisten", "portrait": None}


@router.post("/conversations/{cid}/voice-image")
async def voice_image(cid: str, x: VoiceImageIn, u: dict = Depends(current_user)):
    """Voice-call tool: render an image (optionally editing the latest one) and post it to the chat panel."""
    over = await quota_exceeded(u)
    if over:
        raise HTTPException(402, quota_message(over))
    await rate_limit(u, "generation")
    persona = await _voice_persona(cid, u)
    ref = await _latest_media(cid, "image") if x.edit_previous else None
    try:
        out = await run_image_tool(u["id"], x.prompt, ref["path"] if ref else None, "gemini-image", *norm_image_opts(x.aspect_ratio, x.quality))
    except Exception as exc:
        logging.getLogger("chat").warning("voice image failed: %s", exc)
        raise HTTPException(502, "Gambar belum berhasil dibuat, coba lagi") from exc
    await record_usage(u["id"], "image_generation", out["credits"], {"conversation_id": cid, "persona_id": persona["id"], "via": "realtime"})
    out["media"][0]["request"] = (x.request or x.prompt)[:300]
    msg = await _save_ai_msg(cid, persona, IMAGE_EDIT_DONE_TEXT if ref else IMAGE_DONE_TEXT, out["credits"], "text", {"media": out["media"], "tool": "image_edit" if ref else "image"})
    await notify(cid, {"type": "message", "role": "assistant", "persona_id": persona["id"], "message": clean(msg)})
    return {"ok": True, "credits": out["credits"], "edited": bool(ref), "message_id": msg["id"]}


@router.post("/conversations/{cid}/voice-video")
async def voice_video(cid: str, x: VoiceVideoIn, u: dict = Depends(current_user)):
    """Voice-call tool: post the Seedance 2.0/2.5 choice card (or the Drive/credits notice) to the chat panel; the user taps a model there."""
    persona = await _voice_persona(cid, u)
    text, extra = await video_offer(u, cid, {"video_prompt": x.prompt, "duration": x.duration, "aspect_ratio": x.aspect_ratio, "resolution": x.resolution, "real_person": x.real_person, "with_audio": x.with_audio, "from_image": x.from_image}, x.request or x.prompt)
    msg = await _save_ai_msg(cid, persona, text, 0, "text", {**extra, "tool": extra.get("tool", "video")})
    await notify(cid, {"type": "message", "role": "assistant", "persona_id": persona["id"], "message": clean(msg)})
    pt = extra.get("pending_tool")
    return {"ok": True, "needs_choice": bool(pt), "summary": re.sub(r"[*_]", "", text), "options": [{k: o[k] for k in ("tier", "label", "per_sec", "credits", "available")} for o in (pt or {}).get("options", [])]}


@router.post("/conversations/{cid}/social-publish")
async def conv_social_publish(cid: str, x: SocialChatIn, u: dict = Depends(current_user)):
    """Voice tool: publish the latest media of this conversation and drop the result card into the chat."""
    conv = await db.conversations.find_one({"id": cid}, {"_id": 0})
    if not conv or not _can_access(conv, u):
        raise HTTPException(404, "Conversation not found")
    out = await social_publish_from_chat(cid, u, x.providers, x.caption, x.kind, x.app_url)
    persona = (await _get_personas([conv.get("persona_id") or (conv.get("persona_ids") or [None])[0]]) or [None])[0]
    if persona:
        await _save_ai_msg(cid, persona, out["text"], 0, "text", {"tool": "social_publish"})
    return out


async def _github_turn(ctx, plan: dict):
    """GitHub/GitLab Level 1 from chat: list repos, read tree/file, list issues, or author + open a pull/merge request."""
    tool = plan["tool"]
    prov = _repo_provider(tool.split("_", 1)[0])
    kind = tool.split("_", 1)[1]
    gh_connected, gh_repos, gh_tree, gh_read, gh_issues, gh_create_pr = prov["connected"], prov["repos"], prov["tree"], prov["read"], prov["issues"], prov["create"]
    yield ctx.sse(start=True)
    if not await gh_connected(ctx.user["id"]):
        async for ev in _emit_final(ctx, f"{prov['label']} belum terhubung. Hubungkan dulu di menu [Integrasi](/integrations) (tempel Personal Access Token), lalu minta lagi ya.", 0, {"tool": tool, "error": True}):
            yield ev
        return
    yield ctx.sse(status=f"Menghubungi {prov['label']}...")
    repo = (plan.get("repo") or (ctx.task or {}).get("github_repo") or "").strip()
    credits = 0
    try:
        if kind == "repos":
            rows = await gh_repos(ctx.user["id"], plan.get("query") or "")
            text = ("Ini repositori kamu:\n\n" + "\n".join(f"- [{r['full_name']}]({r['url']}){' 🔒' if r['private'] else ''}{' — ' + r['description'][:80] if r['description'] else ''}" for r in rows)) if rows else "Belum ada repositori yang bisa aku lihat dengan token ini."
            extra = {"tool": tool, "results": rows[:30]}
        elif kind == "read":
            path = (plan.get("path") or "").strip("/")
            if not repo:
                raise HTTPException(400, "Sebutkan nama repo-nya (owner/repo) ya.")
            if not path:
                t = await gh_tree(ctx.user["id"], repo)
                listing = "\n".join(f"- `{p}`" for p in t["paths"][:120])
                text = f"Struktur **{t['repo']}** (branch `{t['ref']}`, {t['count']} file{' — dipotong' if t['truncated'] or t['count'] > 120 else ''}):\n\n{listing}"
            else:
                f = await gh_read(ctx.user["id"], repo, path)
                if f.get("dir"):
                    text = f"Isi folder `{path}` di **{repo}**:\n\n" + "\n".join(f"- {'📁' if e['type'] == 'dir' else '📄'} `{e['path']}`" for e in f["entries"])
                else:
                    yield ctx.sse(status="Membaca & merangkum file...")
                    summary = await llm_text(ctx.system, f"The user asked: {ctx.user_text}\n\nFile `{f['path']}` from repo {repo}:\n```\n{f['content'][:30000]}\n```\n\nAnswer their request about this file in their language (explain/summarize/review as asked). Be concrete and reference line-level details where useful.", ctx.model_key)
                    credits = text_credits(f["content"][:30000], summary, ctx.model_key)
                    text = f"{summary}\n\n[Lihat file di {prov['label']}]({f['url']})"
            extra = {"tool": tool}
        elif kind == "issues":
            if not repo:
                raise HTTPException(400, "Sebutkan nama repo-nya (owner/repo) ya.")
            rows = await gh_issues(ctx.user["id"], repo, plan.get("state") or "open")
            text = (f"Issue/PR **{repo}** ({plan.get('state') or 'open'}):\n\n" + "\n".join(f"- [#{r['number']} {r['title']}]({r['url']}){' (PR)' if r['is_pr'] else ''} — {r['author']}" for r in rows)) if rows else f"Tidak ada issue {plan.get('state') or 'open'} di {repo}."
            extra = {"tool": tool, "results": rows}
        elif kind == "review":
            if not repo:
                raise HTTPException(400, "Sebutkan nama repo-nya (owner/repo) ya.")
            yield ctx.sse(status=f"Mengambil diff {prov['pr'].lower()}...")
            d = await prov["diff"](ctx.user["id"], repo, _pr_number(plan.get("number")))
            yield ctx.sse(status=f"Meninjau {len(d['files'])} file...")
            text, credits = await git_review_text(prov, d, ctx.user, ctx.model_key, ctx.user_text)
            extra = {"tool": tool, "review_of": {"repo": repo, "number": d["number"], "url": d["url"]}}
        else:  # pr | commit
            if not repo:
                raise HTTPException(400, "Sebutkan nama repo-nya (owner/repo) ya.")
            branch = (plan.get("branch") or "").strip() if kind == "commit" else ""
            if kind == "commit" and not branch:
                raise HTTPException(400, "Sebutkan branch tujuan commit-nya ya (mis. `dev`) — atau minta aku buat pull request saja.")
            files = [p.strip("/") for p in (plan.get("files") or []) if p]
            tree = await gh_tree(ctx.user["id"], repo)
            if not files:
                yield ctx.sse(status="Memilih file yang perlu diubah...")
                pick = await llm_json("Pick the repository files that must be edited or created to fulfil the request. Reply JSON {\"files\":[path,...]} (max 6, existing paths from the list or new paths).",
                                      f"Request: {plan.get('instructions') or ctx.user_text}\n\nRepo files:\n" + "\n".join(tree["paths"][:400]))
                files = [str(p).strip("/") for p in (pick.get("files") or [])][:6]
            if not files:
                raise HTTPException(400, "Aku belum bisa menentukan file mana yang harus diubah — sebutkan path file-nya.")
            yield ctx.sse(status=f"Membaca {len(files)} file & menyusun perubahan...")
            current = {}
            for p in files:
                try:
                    f = await gh_read(ctx.user["id"], repo, p)
                    current[p] = f["content"] if not f.get("dir") else ""
                except HTTPException:
                    current[p] = ""
            ctx_files = "\n\n".join(f"=== {p} ({'NEW FILE' if not c else 'current content'}) ===\n{c[:25000]}" for p, c in current.items())
            res = await llm_json(f"You are a senior engineer preparing a {'direct commit' if kind == 'commit' else 'pull request'}. Return the COMPLETE new content of every file you change (no diffs, no placeholders, no truncation). "
                                 "Reply JSON only: {\"title\":str,\"body\":str (markdown summary of changes),\"changes\":[{\"path\":str,\"content\":str}|{\"path\":str,\"delete\":true}]}. Keep unrelated code untouched.",
                                 f"Request: {plan.get('instructions') or ctx.user_text}\nSuggested title: {plan.get('title') or ''}\n\n{ctx_files}", ctx.model_key)
            changes = [c for c in (res.get("changes") or []) if isinstance(c, dict) and c.get("path")]
            credits = text_credits(ctx_files, json.dumps(changes, ensure_ascii=False), ctx.model_key)
            title = (res.get("title") or plan.get("title") or "Perubahan dari Oryntix")[:200]
            if kind == "commit":
                yield ctx.sse(status=f"Commit langsung ke branch `{branch}`...")
                cm = await prov["commit"](ctx.user["id"], repo, title, changes, branch)
                note = " (branch baru dibuat dari default)" if cm.get("created_branch") else (" — ini default branch" if cm["branch"] == cm.get("default_branch") else "")
                text = f"Commit **{cm['title']}** sudah masuk ke branch `{cm['branch']}`{note} di **{repo}** — [lihat di {prov['label']}]({cm['url']}).\n\nFile: " + ", ".join(f"`{p}`" for p in cm["files"]) + f"\n\n{res.get('body') or ''}"
                extra = {"tool": tool, "github_commit": cm}
            else:
                yield ctx.sse(status="Membuat branch, commit & pull request...")
                pr = await gh_create_pr(ctx.user["id"], repo, title, res.get("body") or "", changes)
                text = f"{prov['pr']} **!{pr['number']} {pr['title']}** sudah dibuka di **{repo}** — [lihat di {prov['label']}]({pr['url']}).\n\nBranch `{pr['branch']}` → `{pr['base']}` · file: " + ", ".join(f"`{p}`" for p in pr["files"]) + f"\n\n{res.get('body') or ''}"
                extra = {"tool": tool, "github_pr": pr}
    except HTTPException as e:
        text, extra = f"{e.detail}", {"tool": tool, "error": True}
    if credits:
        await record_usage(ctx.user["id"], "chat", credits, {"conversation_id": ctx.cid, "persona_id": ctx.persona["id"], "model": ctx.model_key, "tool": tool})
    async for ev in _emit_final(ctx, text, credits, extra):
        yield ev


async def _tool_turn(ctx, plan: dict):
    """Create an image/document from chat, or run a Google Drive / GitHub action. Expensive tools (≥ threshold) ask for confirmation first."""
    if plan["tool"].startswith("drive_"):
        gen = _drive_turn(ctx, plan)
    elif plan["tool"].startswith(("github_", "gitlab_")):
        gen = _github_turn(ctx, plan)
    elif plan["tool"] == "social_publish":
        gen = _social_turn(ctx, plan)
    else:
        cfg = await get_routing()
        gen = {"image": _image_turn, "image_edit": _image_turn, "video": _video_turn}.get(plan["tool"], lambda c, p, _t: _document_turn(c, p))(ctx, plan, cfg["confirm_threshold"])
    async for ev in gen:
        yield ev


@dataclass
class ReplyCtx:
    """Everything one persona needs to answer a turn."""
    cid: str
    user: dict
    persona: dict
    roster: Optional[list]
    prompt: str
    voice_mode: bool = False
    via: Optional[str] = None
    user_text: str = ""
    attach_len: int = 0
    meta_extra: dict = field(default_factory=dict)
    system: str = ""
    model_key: Optional[str] = None
    routed: Optional[str] = None
    model_choice: Optional[str] = None
    choice_needed: Optional[dict] = None
    task: Optional[dict] = None

    @property
    def meta(self) -> dict:
        p = self.persona
        m = {"persona_id": p["id"], "persona_name": p["name"], "portrait": p.get("portrait"), "voice": p.get("voice", "alloy")}
        if self.routed:
            m.update(model_label=model_label(self.model_key), routed=self.routed)
        return m

    def sse(self, **payload) -> str:
        return f"data: {json.dumps({**self.meta, **payload})}\n\n"


async def _prepare_ctx(ctx: ReplyCtx) -> ReplyCtx:
    ctx.system = await _persona_system(ctx.persona, ctx.user, ctx.roster, voice_mode=ctx.voice_mode, query=ctx.user_text)
    if ctx.via == "meeting_chat":
        ctx.system += "\n\n" + MEETING_CHAT_STYLE
    own = ctx.persona.get("model")
    suggested, reason = await route_model(own, ctx.user_text, ctx.attach_len, await _owner_settings(ctx.user))
    if ctx.model_choice:  # the user already answered the confirmation card
        ctx.model_key = own if ctx.model_choice == "__own__" else ctx.model_choice
        ctx.routed = reason if ctx.model_key != own else None
    elif reason and suggested != own and not ctx.voice_mode and ctx.via != "meeting_chat":
        # never switch silently: ask first (default = the assistant's own model). Video/Seedance flows are separate and untouched.
        ctx.model_key, ctx.routed = own, None
        ctx.choice_needed = {"reason": reason, "options": [{"id": "__own__", "label": model_label(own), "default": True}, {"id": suggested, "label": model_label(suggested), "default": False}]}
    else:
        ctx.model_key, ctx.routed = (suggested, reason) if ctx.voice_mode or ctx.via == "meeting_chat" else (own, None)
    if ctx.routed:  # another model answers this turn — restate the language rule so the switch is invisible to the user
        ctx.system += f"\n\nMODEL SWITCH NOTE: you are a different model handling this turn on behalf of the same persona. Keep the same persona, tone and language. {lang_rule(ctx.user)}"
    last = await db.messages.find_one({"conversation_id": ctx.cid, "role": "assistant", "tool": "workspace_search"}, {"_id": 0, "results": 1}, sort=[("created_at", -1)])
    if last and last.get("results"):
        top = await db.tasks.find_one({"id": last["results"][0]["id"]}, {"_id": 0, "goal": 1, "final_output": 1})
        if top:
            ctx.system += f"\n\nWORKSPACE ITEM RECENTLY FOUND (you may quote it; link: /workspace/{last['results'][0]['id']}):\nTitle: {top.get('goal')}\n{(top.get('final_output') or '')[:3000]}"
    conv = await db.conversations.find_one({"id": ctx.cid}, {"_id": 0, "task_id": 1}) or {}
    if conv.get("task_id"):
        ctx.task = await db.tasks.find_one({"id": conv["task_id"]}, {"_id": 0})
        if ctx.task:
            ctx.system += TASK_CONTEXT.format(tid=ctx.task["id"], ver=ctx.task.get("version") or 1, status=ctx.task.get("status"), goal=ctx.task.get("goal"), body=(ctx.task.get("final_output") or "(belum ada hasil)")[:24000])
    return ctx


DELEGATE_RE = re.compile(r"\b(terima beres|beres saja|kerjakan saja|langsung (saja|kerjakan)|serahkan|tolong kerjakan)\b", re.I)
DISCUSS_RE = re.compile(r"\b(satu per satu|bahas (dulu|bersama|saja)|diskusi(kan)? dulu|pelan-pelan|bertahap)\b", re.I)
TASK_RE_STRONG = re.compile(r"\b(buatkan|susun(kan)?|kerjakan|siapkan|rancang|tulis(kan)?)\b", re.I)
TEAM_RE = re.compile(r"\b(bagi(kan)? (tugas(nya)? )?ke tim|bagi tugas|delegasikan|kerjakan bersama tim|libatkan (tim|asisten lain)|split to team)\b", re.I)
SEARCH_RE = re.compile(r"\b(cari(kan)?|carilah|temukan|ada (dokumen|hasil|notulen|laporan|file|berkas|tugas)|dokumen (tentang|mengenai|soal)|di ruang kerja|workspace)\b", re.I)
WEB_RE = re.compile(r"\b(web|internet|google|online|daring|berita|terkini|terbaru|hari ini|saat ini|sekarang)\b", re.I)  # "cari di web…" is for the model's search tool, not the workspace
TEAM_OFFER = ("\n\nAtau, karena tim kita ada beberapa asisten, saya juga bisa **bagi ke tim**: saya pecah jadi sub-tugas untuk asisten yang paling cocok, lalu saya rangkai hasilnya. "
              "Anda juga boleh menunjuk langsung, misalnya «bagian keuangan minta Nova».")


async def _search_turn(ctx):
    from assignments import search_workspace, results_markdown
    yield ctx.sse(start=True)
    yield ctx.sse(status="Mencari di Ruang Kerja...")
    results = await search_workspace(ctx.user, ctx.user_text)
    async for ev in _emit_final(ctx, results_markdown(results, ctx.user_text[:80]), 0, {"tool": "workspace_search", "results": results}):
        yield ev


async def _calendar_turn(ctx):
    """'Catat rapat besok jam 10, ingatkan 1 jam sebelum lewat chat' → event (+ linked reminder) or a clarifying question."""
    from tools import plan_calendar
    from assignments import EventIn, create_event_doc, event_markdown
    tz = (ctx.user.get("settings") or {}).get("timezone")
    conv = await db.conversations.find_one({"id": ctx.cid}, {"_id": 0, "pending_calendar": 1}) or {}
    pending = conv.get("pending_calendar")
    text = f"{pending['text']}\nJawaban user atas pertanyaan «{pending['question']}»: {ctx.user_text}" if pending else ctx.user_text
    plan = await plan_calendar(text, tz, ctx.prompt[-600:] if ctx.prompt else "", force=bool(pending))
    if pending:
        await db.conversations.update_one({"id": ctx.cid}, {"$set": {"pending_calendar": None}})
    if not plan:
        return
    yield ctx.sse(start=True)
    if not plan.get("start_at") or plan.get("question"):
        q = plan.get("question") or f"Siap, saya catat «{plan.get('title') or 'kegiatan ini'}». Tanggal dan jam berapa tepatnya?"
        await db.conversations.update_one({"id": ctx.cid}, {"$set": {"pending_calendar": {"text": text[:1500], "question": q[:300], "at": now_iso()}}})
        async for ev in _emit_final(ctx, q, 0, {"tool": "calendar_question"}):
            yield ev
        return
    offsets = [int(o) for o in (plan.get("remind_offsets") or []) if isinstance(o, (int, float))]
    no_reminder = offsets == [0]
    offsets = [o for o in offsets if o > 0] or ([] if no_reminder else [30])
    mode = plan.get("remind_mode") if plan.get("remind_mode") in ("call", "chat") else "call"
    try:
        ev = await create_event_doc(ctx.user, EventIn(title=(plan.get("title") or ctx.user_text[:80]).strip()[:200], start_at=plan["start_at"], notes=(plan.get("notes") or "")[:2000],
                                                      remind_mode=mode if offsets else None, remind_offsets=offsets, persona_id=ctx.persona["id"],
                                                      repeat=plan.get("repeat") if plan.get("repeat") in ("daily", "weekly", "monthly") else "none"))
    except Exception:
        async for e in _emit_final(ctx, "Maaf, waktu yang disebut belum bisa saya pahami. Bisa sebutkan tanggal dan jamnya?", 0, {"tool": "calendar_question"}):
            yield e
        return
    async for e in _emit_final(ctx, event_markdown(ev), 0, {"tool": "calendar_event", "event": ev, "cta": {"label": "Buka Kalender", "href": "/calendar"}}):
        yield e


async def _task_offer_turn(ctx):
    """Long/scheduled delegation → offer 'bahas satu per satu' vs 'terima beres'; a reply to a pending offer is resolved here."""
    conv = await db.conversations.find_one({"id": ctx.cid}, {"_id": 0, "pending_task": 1, "persona_ids": 1}) or {}
    if (conv.get("persona_ids") or [ctx.persona["id"]])[0] != ctx.persona["id"]:
        return  # in group chats only the first assistant handles task offers (avoids duplicate offers)
    if conv.get("pending_task"):
        resolved = False
        async for ev in _resolve_pending_offer(ctx, conv):
            resolved = True
            yield ev
        if resolved:
            return
        await db.conversations.update_one({"id": ctx.cid}, {"$set": {"pending_task": None}})  # user moved on
    plan = await plan_task(ctx.user_text, (ctx.user.get("settings") or {}).get("timezone"), ctx.prompt[-600:] if ctx.prompt else "")
    if not plan.get("is_task") or not (plan.get("long") or plan.get("scheduled_at")):
        return
    async for ev in _make_offer(ctx, plan):
        yield ev


async def _offer_mode(ctx, pending: dict) -> Optional[str]:
    """How the user answered a pending task offer: team / delegate / discuss, or None when the reply is unrelated."""
    from assignments import names_mentioned
    named = []
    if pending.get("team_possible"):
        roster = await db.personas.find({"user_id": workspace_id(ctx.user), "deleted": {"$ne": True}}, {"_id": 0, "id": 1, "name": 1}).to_list(50)
        named = names_mentioned(ctx.user_text, roster, exclude_id=ctx.persona["id"])
    if TEAM_RE.search(ctx.user_text) or named:
        return "team"
    if DELEGATE_RE.search(ctx.user_text):
        return "delegate"
    return "discuss" if DISCUSS_RE.search(ctx.user_text) else None


async def _resolve_pending_offer(ctx, conv: dict):
    from assignments import accept_pending
    mode = await _offer_mode(ctx, conv["pending_task"])
    if not mode:
        return
    yield ctx.sse(start=True)
    yield ctx.sse(status={"delegate": "Mencatat tugas...", "team": "Membagi tugas ke tim..."}.get(mode, "Menyusun langkah..."))
    pending = {**conv["pending_task"], "directives": ((conv["pending_task"].get("directives") or "") + "\n" + ctx.user_text).strip()}
    msg = await accept_pending({"id": ctx.cid, **conv, "pending_task": pending, "persona_id": ctx.persona["id"]}, ctx.user, mode)
    yield ctx.sse(final=True, content=msg["content"], message_id=msg["id"], credits_used=msg.get("credits", 0), **{k: msg[k] for k in ("tool", "task_id") if k in msg})


async def _make_offer(ctx, plan: dict):
    from assignments import offer_text
    team_possible = bool(plan.get("long")) and await db.personas.count_documents({"user_id": workspace_id(ctx.user), "deleted": {"$ne": True}}) > 1
    await db.conversations.update_one({"id": ctx.cid}, {"$set": {"pending_task": {**plan, "persona_id": ctx.persona["id"], "team_possible": team_possible, "directives": ctx.user_text, "offered_at": now_iso()}}})
    yield ctx.sse(start=True)
    async for ev in _emit_final(ctx, offer_text(plan, ctx.user) + (TEAM_OFFER if team_possible else ""), 0,
                                {"tool": "task_offer", "pending_task": {"title": plan.get("title"), "scheduled_at": plan.get("scheduled_at"), "team_possible": team_possible}}):
        yield ev


async def _wants_revision(text: str, task: dict, history: str = "") -> bool:
    if not REVISE_RE.search(text or ""):
        return False
    try:
        r = await llm_json("Does the user's LAST message ask to CHANGE/REVISE/UPDATE the workspace document under discussion — including applying, saving or "
                           "using content that the assistant proposed earlier in the conversation as the new document content? Questions, opinions, asking for a "
                           "proposal/draft in chat only, or chit-chat are NOT revisions. Reply JSON {\"revise\": true|false}.",
                           f"Document title: {task.get('goal')}\nRecent conversation:\n{history}\n\nUser LAST message: {text}")
        return bool(r.get("revise"))
    except Exception:
        return False


async def _revise_turn(ctx):
    yield ctx.sse(start=True)
    yield ctx.sse(status="Merevisi hasil tugas di Ruang Kerja...")
    try:
        new_md, summary, used = await revise_with_llm(ctx.task, ctx.user_text, ctx.system, ctx.model_key, await _history_text(ctx.cid, limit=12, with_summary=False))
    except Exception:
        async for ev in _emit_final(ctx, "Maaf, revisinya belum berhasil disimpan. Coba ulangi permintaannya.", 0, {}):
            yield ev
        return
    ver = await save_revision(ctx.task, new_md, summary or ctx.user_text, ctx.persona)
    await record_usage(ctx.user["id"], "task_revision", used, {"task_id": ctx.task["id"], "conversation_id": ctx.cid, "actor_id": ctx.user["id"]})
    text = f"Revisi **v{ver}** tersimpan di Ruang Kerja ✅\n\n{summary or 'Perubahan sesuai permintaan Anda sudah diterapkan.'}\n\n[Lihat hasil terbaru](/workspace/{ctx.task['id']})"
    async for ev in _emit_final(ctx, text, used, {"tool": "revise", "task_id": ctx.task["id"], "task_version": ver}):
        yield ev


async def _typed_intercepts(ctx: ReplyCtx):
    """Special handling for a TYPED user message (archive recall, web search, task offer, revision, tools). Yields nothing when none applies."""
    from archives import archive_turn_text
    hit = await archive_turn_text(ctx)
    if hit:
        yield ctx.sse(start=True)
        async for ev in _emit_final(ctx, hit["content"], 0, hit["extra"]):
            yield ev
        return
    if SEARCH_RE.search(ctx.user_text) and not TASK_RE_STRONG.search(ctx.user_text) and not WEB_RE.search(ctx.user_text) and not CAL_RE.search(ctx.user_text):
        async for ev in _search_turn(ctx):
            yield ev
        return
    if CAL_RE.search(ctx.user_text) or await db.conversations.count_documents({"id": ctx.cid, "pending_calendar": {"$ne": None}}):
        hit = False
        async for ev in _calendar_turn(ctx):
            hit = True
            yield ev
        if hit:
            return
    plan = None
    if GITHUB_RE.search(ctx.user_text) or SOCIAL_RE.search(ctx.user_text):
        # repo/PR/MR and social-media requests are tool calls, never a "big task" offer
        plan = await plan_tool(ctx.user_text, ctx.prompt, (ctx.user.get('settings') or {}).get('timezone'))
        if plan.get("tool", "none").startswith(("github_", "gitlab_", "social_")):
            async for ev in _tool_turn(ctx, plan):
                yield ev
            return
    if not ctx.task:
        handled = False
        async for ev in _task_offer_turn(ctx):
            handled = True
            yield ev
        if handled:
            return
    elif await _wants_revision(ctx.user_text, ctx.task, await _history_text(ctx.cid, limit=8, with_summary=False)):
        async for ev in _revise_turn(ctx):
            yield ev
        return
    if not ctx.persona.get("builtin") and (wants_tool(ctx.user_text) or ((EDIT_RE.search(ctx.user_text) or ANIMATE_RE.search(ctx.user_text)) and await _latest_media(ctx.cid, "image"))):
        plan = plan or await plan_tool(ctx.user_text, ctx.prompt, (ctx.user.get('settings') or {}).get('timezone'))
        if plan.get("tool") != "none":
            async for ev in _tool_turn(ctx, plan):
                yield ev


async def _plain_reply(ctx: ReplyCtx):
    """Default streamed LLM answer; the last yielded item is the int credits used."""
    yield ctx.sse(start=True)
    from provider_tools import persona_tools, run_with_tools
    tool_ids = persona_tools(ctx.persona, ctx.model_key)
    tool_out = None
    if tool_ids:
        yield ctx.sse(status="Menggunakan alat bawaan model…")
        try:
            tool_out = await run_with_tools(ctx.system, ctx.prompt, ctx.model_key, tool_ids, ctx.user["id"], [ctx.persona["vector_store_id"]] if ctx.persona.get("vector_store_id") else None)
        except Exception as exc:
            logging.getLogger(__name__).warning("provider tools failed, falling back: %s", str(exc)[:200])
    if tool_out and (tool_out["text"] or tool_out["media"] or tool_out.get("pending_files")):
        full = tool_out["text"] or "Ini hasilnya 🎨"
    else:
        try:
            from pricing import get_pricing as _gp, model_price, affordable_tokens
            from llm import owner_balance_exact
            pp = await _gp()
            mp = model_price(pp, ctx.model_key)
            cap = affordable_tokens(pp, await owner_balance_exact(ctx.user), mp["out"], "text") if mp else None
            full = await llm_text(ctx.system, ctx.prompt, ctx.model_key, max_tokens=cap)
            if cap is not None and not full.strip():
                full = "Maaf, sisa kredit workspace terlalu sedikit untuk menghasilkan jawaban lengkap. Silakan isi ulang kredit terlebih dahulu."
        except Exception:
            full = "Maaf, terjadi gangguan saat menghasilkan jawaban. Silakan coba lagi."
    if not ctx.voice_mode and not tool_out and not ctx.persona.get("builtin") and is_media_refusal(ctx.user_text, full):
        # the model wrongly claimed it cannot render — render it anyway
        plan = await plan_tool(ctx.user_text, ctx.prompt, (ctx.user.get('settings') or {}).get('timezone'))
        if plan.get("tool") not in ("image", "image_edit", "video"):
            plan = {"tool": "video" if re.search(r"\b(video\w*|klip|clip|animasi|reels?)\b", ctx.user_text, re.I) else "image", "image_prompts": [ctx.user_text], "image_prompt": ctx.user_text, "video_prompt": ctx.user_text}
        async for ev in _tool_turn(ctx, plan):
            yield ev
        return
    yield ctx.sse(delta=full)  # whole reply at once — no word-by-word typing effect
    used = text_credits(ctx.prompt, full, ctx.model_key)
    await record_usage(ctx.user["id"], "chat", used, {"conversation_id": ctx.cid, "persona_id": ctx.persona["id"], "model": ctx.model_key, **ctx.meta_extra})
    extra = {"model_key": ctx.model_key, "model_label": model_label(ctx.model_key), "routed": ctx.routed} if ctx.routed else {}
    if tool_out and tool_out.get("tools_used"):
        for t in tool_out["tools_used"]:
            if t["credits"]:
                await record_usage(ctx.user["id"], f"tool:{t['id']}", t["credits"], {"conversation_id": ctx.cid, "persona_id": ctx.persona["id"], "count": t["count"], **ctx.meta_extra})
        used += tool_out["tool_credits"]
        extra.update({"tools_used": tool_out["tools_used"], "citations": tool_out["citations"]})
    if tool_out and tool_out.get("media"):
        extra["media"] = tool_out["media"]
    if tool_out and tool_out.get("pending_files"):
        extra["pending_files"] = tool_out["pending_files"]
    ai_msg = await _save_ai_msg(ctx.cid, ctx.persona, full, used, ctx.via, extra)
    yield ctx.sse(final=True, message_id=ai_msg["id"], content=full)
    await notify(ctx.cid, {"type": "message", "role": "assistant", "persona_id": ctx.persona["id"], "message": clean(ai_msg)})
    yield used


async def _persona_reply(ctx: ReplyCtx):
    """Stream one persona reply as SSE strings; the last yielded item is the int credits used."""
    await _prepare_ctx(ctx)
    if ctx.choice_needed:
        rs = {"it": "pertanyaan kode/IT", "research": "riset atau konteks panjang"}.get(ctx.choice_needed["reason"], ctx.choice_needed["reason"])
        text = f"Untuk {rs} ini saya bisa minta bantuan model lain. Mau dikerjakan dengan model mana?"
        async for ev in _emit_final(ctx, text, 0, {"tool": "model_choice", "choice": {**ctx.choice_needed, "user_text": ctx.user_text}}):
            yield ev
        return
    if ctx.user_text and not ctx.voice_mode:
        intercepted = False
        async for ev in _typed_intercepts(ctx):
            intercepted = True
            yield ev
        if intercepted:
            return
    async for ev in _plain_reply(ctx):
        yield ev


async def _finish_stream(cid: str, u: dict, total: int, last_message: str):
    bal = ((await db.users.find_one({"id": workspace_id(u)}, {"_id": 0, "credits": 1})) or {}).get("credits", 0)
    await db.conversations.update_one({"id": cid}, {"$set": {"updated_at": now_iso(), "last_message": last_message[:120], "last_sender_id": u["id"], f"read_at.{u['id']}": now_iso()}})
    return f"data: {json.dumps({'done': True, 'credits_used': total, 'credits': bal})}\n\ndata: [DONE]\n\n"


def _sse(gen):
    return StreamingResponse(gen, media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


async def _store_user_message(cid: str, x: MsgIn, u: dict, attach_text: str, attach_meta: list) -> dict:
    user_msg = {"id": new_id(), "conversation_id": cid, "role": "user", "content": x.content,
                "attachments": attach_meta, "sender_user_id": u["id"], "sender_name": u.get("name") or "User",
                "created_at": now_iso()}
    if x.channel:
        user_msg["via"] = x.channel
    if x.reply_to and x.reply_to.get("content"):
        user_msg["reply_to"] = {"id": str(x.reply_to.get("id") or ""), "name": str(x.reply_to.get("name") or "")[:80], "content": str(x.reply_to["content"])[:400]}
    if x.forwarded:
        user_msg["forwarded"] = True
    if attach_text:
        user_msg["attachment_text"] = attach_text[:4000]
    await db.messages.insert_one(dict(user_msg))
    await notify(cid, {"type": "message", "role": "user", "sender_name": user_msg["sender_name"], "message": clean(user_msg)})
    await _fanout_message(cid, u["id"], user_msg["sender_name"], x.content)
    return user_msg


async def _fanout_message(cid: str, sender_id: str, sender_name: str, content: str):
    """Per-user event (sidebar refresh) + push for human participants of a group/friend chat who are not the sender."""
    from realtime import notify_users
    from push import send_push
    conv = await db.conversations.find_one({"id": cid}, {"_id": 0, "participants": 1, "user_id": 1, "title": 1, "type": 1})
    if not conv:
        return
    others = [p for p in set((conv.get("participants") or []) + [conv.get("user_id")]) if p and p != sender_id]
    # bump the row first so the conversation snapshot riding on the event is already current
    await db.conversations.update_one({"id": cid}, {"$set": {"updated_at": now_iso(), "last_message": content[:120], "last_sender_id": sender_id}})
    await notify_users(others, {"type": "message_new", "conversation_id": cid, "sender_name": sender_name, "preview": content[:80]})
    for o in others:
        await send_push(o, f"{sender_name} · {conv.get('title') or 'Oryntix'}", content[:120], {"link": f"/chat/{cid}", "tag": f"conv-{cid}"}, kind="messages")


def _pick_responders(x: MsgIn, personas: list) -> list:
    mentioned = _mentioned(x.content, personas)
    if mentioned:
        return mentioned
    return personas[:1] if x.channel == "meeting_chat" else personas  # text side-channel: one assistant answers


AI_ADDRESS_RE = re.compile(r"\b(tolong|bisa(kah)?|buatkan|carikan|jelaskan|rangkum|ringkas|analisis|hitung|terjemahkan|asisten|ai\b|bot|menurut(mu)?|bantu|please|can you|could you)\b|\?\s*$", re.I)


def _addressed_to_ai(text: str, personas: list) -> bool:
    low = (text or "").lower()
    return any(p["name"].lower() in low for p in personas) or bool(AI_ADDRESS_RE.search(text or ""))


async def _route_group(x: MsgIn, personas: list, cid: str) -> list:
    """Group chat without @mention: let only the 1-2 most relevant assistants answer (saves N× context tokens)."""
    text = (x.content or "").strip()
    last = await db.messages.find_one({"conversation_id": cid, "role": "assistant", "archived": {"$ne": True}}, {"_id": 0, "persona_id": 1}, sort=[("created_at", -1)])
    if len(text) < 25 and last:  # short follow-up ("ya", "lanjutkan") → whoever spoke last
        return [p for p in personas if p["id"] == last.get("persona_id")] or personas[:1]
    if re.search(r"\b(semua|kalian|masing-masing|everyone|all of you|pendapat kalian)\b", text, re.I):
        return personas
    roster = "\n".join(f"- {p['name']}: {(p.get('summary') or (p.get('profile') or {}).get('system_instructions') or '')[:160]}" for p in personas)
    try:
        r = await llm_json("Pick which assistants (1, at most 2) should answer the user's latest group message, based on relevance to their role. Reply JSON {\"names\": [str]}.",
                           f"Assistants:\n{roster}\nUser message: {text[:600]}")
        names = {str(n).lower() for n in (r.get("names") or [])}
        chosen = [p for p in personas if p["name"].lower() in names][:2]
        return chosen or personas[:1]
    except Exception:
        return personas[:1]


def _reply_extra(x: MsgIn, attach_text: str) -> str:
    notes = [f"\n\n[Lampiran dari user]:\n{attach_text}"] if attach_text else []
    if x.reply_to and x.reply_to.get("content"):
        notes.append(f"\n[User membalas (quote) pesan dari {x.reply_to.get('name') or 'asisten'}: \"{str(x.reply_to['content'])[:400]}\" — tanggapi dalam konteks kutipan itu.]")
    if x.forwarded:
        notes.append("\n[Pesan ini DITERUSKAN user dari percakapan lain; tanggapi isinya.]")
    if x.interrupted:
        notes.append("\n[Catatan: user baru saja menyela saat asisten sedang berbicara. Tanggapi langsung apa yang user katakan.]")
    if x.channel == "meeting_chat":
        notes.append("\n[Catatan: pesan terakhir user DIKETIK di panel chat panggilan; jawab dalam bentuk teks/markdown.]")
    return "".join(notes)


async def _collect(gen, totals: list):
    """Re-yield SSE strings from a reply generator; its trailing int (credits) is added to totals[0]."""
    async for ev in gen:
        if isinstance(ev, (int, float)) and not isinstance(ev, bool):
            totals[0] += ev
        else:
            yield ev


async def _moderator_if_stuck(cid: str, conv: dict, responders: list, roster, x: MsgIn):
    """Meeting only: the Moderator steps in when the discussion is stuck (never for the text side-channel)."""
    if conv.get("type") != "meeting" or len(responders) < 2 or x.channel:
        return None
    user_turns = await db.messages.count_documents({"conversation_id": cid, "role": "user"})
    return roster if user_turns >= 2 and await _is_stuck(cid) else None


async def _choose_responders(x: MsgIn, conv: dict, personas: list, cid: str) -> list:
    """Which assistants answer this message: @mentions win; otherwise route by relevance (and only when addressed, in human chats)."""
    if not personas:
        return []
    if _mentioned(x.content, personas):
        return _pick_responders(x, personas)
    humans_chat = len(conv.get("participants") or []) > 1 and conv.get("type") != "meeting"
    if humans_chat:  # people talking to each other: assistants only step in when clearly addressed
        return (await _route_group(x, personas, cid)) if _addressed_to_ai(x.content, personas) else []
    if conv.get("type") == "group" and len(personas) > 1 and x.channel != "meeting_chat":
        return await _route_group(x, personas, cid)
    return _pick_responders(x, personas)


@router.post("/conversations/{cid}/send")
async def send_message(cid: str, x: MsgIn, u: dict = Depends(current_user)):
    await throttle_message(u, x.content)
    conv, personas = await _load_ai_conv(cid, u)
    attach_text, attach_meta = await _process_attachments(x.attachments, u["id"])
    if x.replay and x.model_choice:
        if x.choice_msg_id:
            await db.messages.delete_one({"id": x.choice_msg_id, "conversation_id": cid, "tool": "model_choice"})
    else:
        await _store_user_message(cid, x, u, attach_text, attach_meta)
    responders = await _choose_responders(x, conv, personas, cid)
    if len(responders) > 1 and (GITHUB_RE.search(x.content or "") or SOCIAL_RE.search(x.content or "")):
        responders = responders[:1]  # tool requests (PR/MR, social posting) run once — not one card per persona
    roster = [p["name"] for p in personas] if len(personas) > 1 else None
    extra = _reply_extra(x, attach_text)
    async def _bill_user(p: dict) -> dict:  # each assistant's replies are paid by the assistant's owner
        if p.get("user_id") in (u["id"], u.get("owner_id")):
            return u
        return await db.users.find_one({"id": p.get("user_id")}, {"_id": 0}) or u

    async def stream():
        totals = [0]
        inflight_start(u["id"])
        try:
            if attach_text and x.channel == "meeting_chat":  # let the voice agents hear what was attached in the call chat panel
                yield f"data: {json.dumps({'attachments_context': attach_text[:2000], 'attachment_names': [a.get('name') for a in attach_meta]})}\n\n"
            for persona in responders:
                prompt = (await _history_text(cid, query=x.content)) + extra + f"\n{persona['name']}:"
                ctx = ReplyCtx(cid=cid, user=await _bill_user(persona), persona=persona, roster=roster, prompt=prompt, voice_mode=x.voice_mode,
                               via=x.channel, user_text=x.content, attach_len=len(attach_text), model_choice=x.model_choice)
                async for ev in _collect(_persona_reply(ctx), totals):
                    yield ev
            if await _moderator_if_stuck(cid, conv, responders, roster, x):
                async for ev in _collect(_moderator_interject(cid, u, roster, "stuck"), totals):
                    yield ev
            if await _long_chat(cid, await db.conversations.find_one({"id": cid}, {"_id": 0, "summary_snoozed_at_count": 1})):
                yield f"data: {json.dumps({'summary_request': True})}\n\n"
            yield await _finish_stream(cid, u, totals[0], x.content)
        finally:
            inflight_end(u["id"])

    return _sse(stream())


MOD_META = {"persona_id": "__moderator__", "persona_name": "Moderator", "portrait": None, "is_moderator": True, "moderator_kind": "interject", "voice": "onyx"}


async def _is_stuck(cid: str) -> bool:
    """Cheap LLM check: are the assistants disagreeing without resolution or going in circles?"""
    history = await _history_text(cid, limit=10)
    sys = ("You are a silent meeting observer. Decide if the discussion is STUCK: participants clearly disagree without "
           "converging, keep repeating the same points, or talk past the user's question. Answer with exactly one word: YES or NO.")
    ans = ""
    try:
        ans = (await llm_text(sys, f"Diskusi:\n{history}\n\nStuck?")).strip().upper()
    except Exception:
        return False
    return ans.startswith("YES") or ans.startswith("YA")


async def _moderator_text(cid: str, u: dict, roster: list, reason: str):
    """Generate + persist a short Moderator interjection. Returns (text, credits_used); text is '' when nothing to say."""
    history = await _history_text(cid, limit=12)
    if reason == "silence":
        goal = ("The user has been quiet for a while. In 1-2 warm, short spoken sentences: briefly note what has been agreed so far "
                "(one clause), then gently ask the user whether they want to continue, decide, or wrap up.")
    else:
        goal = ("The discussion is stuck (disagreement or going in circles). In 1-2 short spoken sentences: name the core disagreement "
                "neutrally, then propose a concrete way forward or ask the user to decide. Do NOT summarize everything.")
    sys = (f"You are the meeting Moderator facilitating a LIVE spoken discussion. You MUST write entirely in {_lang_name(u)}. "
           f"{goal} Plain spoken text, no markdown, friendly and natural.")
    try:
        inter = await llm_text(sys, f"Peserta: {', '.join(roster)}\n\nDiskusi terakhir:\n{history}\n\nModerator (singkat):")
    except Exception:
        inter = ""
    if not inter.strip():
        return "", 0
    used = text_credits(history, inter)
    await record_usage(u["id"], "meeting_moderation", used, {"conversation_id": cid, "reason": reason})
    imsg = {"id": new_id(), "conversation_id": cid, "role": "assistant", "content": inter,
            "persona_id": "__moderator__", "persona_name": "Moderator", "is_moderator": True,
            "portrait": None, "credits": used, "created_at": now_iso()}
    await db.messages.insert_one(dict(imsg))
    await notify(cid, {"type": "message", "role": "assistant", "persona_id": "__moderator__", "message": clean(imsg)})
    return inter, used


async def _moderator_interject(cid: str, u: dict, roster: list, reason: str):
    """Yield SSE events for a short Moderator interjection; the final yielded item is the int credits used."""
    inter, used = await _moderator_text(cid, u, roster, reason)
    if not inter:
        yield 0
        return
    yield f"data: {json.dumps({**MOD_META, 'start': True})}\n\n"
    yield f"data: {json.dumps({**MOD_META, 'delta': inter})}\n\n"
    yield f"data: {json.dumps({**MOD_META, 'final': True, 'content': inter})}\n\n"
    yield used


class ModerateIn(BaseModel):
    reason: str = Field(pattern="^(silence|stuck)$")


@router.post("/conversations/{cid}/moderate")
async def moderate(cid: str, x: ModerateIn, u: dict = Depends(current_user)):
    """Realtime meeting: Moderator speaks only when stuck (LLM check) or after long silence. Returns text (or null)."""
    await rate_limit(u, "chat")
    conv, personas = await _load_ai_conv(cid, u)
    if conv.get("type") == "private":
        raise HTTPException(400, "Moderator hanya untuk panggilan")
    user_turns = await db.messages.count_documents({"conversation_id": cid, "role": "user"})
    needed = 2 if x.reason == "stuck" else 1
    if len(personas) < 2 or user_turns < needed or (x.reason == "stuck" and not await _is_stuck(cid)):
        return {"content": None, "voice": "onyx"}
    text, used = await _moderator_text(cid, u, [p["name"] for p in personas], x.reason)
    return {"content": text or None, "voice": "onyx", "credits_used": used}


@router.post("/conversations/{cid}/nudge")
async def nudge(cid: str, u: dict = Depends(current_user)):
    """Silence in a live meeting/call: Moderator (meeting) or the persona (private) gently checks in. Streams SSE."""
    await rate_limit(u, "chat")
    conv, personas = await _load_ai_conv(cid, u)
    history = await _history_text(cid, limit=12)
    if not history.strip():
        raise HTTPException(400, "Belum ada percakapan")
    roster = [p["name"] for p in personas]
    use_mod = conv.get("type") == "meeting" and len(personas) > 1
    persona = personas[0]
    prompt = (f"{history}\n[Catatan: user terdiam cukup lama. Sapa dengan hangat dalam satu kalimat pendek: tanyakan apakah masih ada, "
              f"atau tawarkan bantuan lanjutan. Jangan mengulang jawaban sebelumnya.]\n{persona['name']}:")

    async def stream():
        totals = [0]
        ctx = ReplyCtx(cid=cid, user=u, persona=persona, roster=roster if len(personas) > 1 else None, prompt=prompt, voice_mode=True, meta_extra={"reason": "silence"})
        gen = _moderator_interject(cid, u, roster, "silence") if use_mod else _persona_reply(ctx)
        async for ev in _collect(gen, totals):
            yield ev
        yield f"data: {json.dumps({'done': True, 'credits_used': totals[0]})}\n\ndata: [DONE]\n\n"

    return _sse(stream())


@router.post("/conversations/{cid}/summary")
async def meeting_summary(cid: str, u: dict = Depends(current_user)):
    """Generate a closing Moderator summary, save it to the conversation and as a Workspace notulen."""
    conv = await db.conversations.find_one({"id": cid})
    if not _can_access(conv, u):
        raise HTTPException(404, "Conversation not found")
    if conv.get("type") == "private":
        raise HTTPException(400, "Notulen hanya untuk percakapan grup atau panggilan")
    history = await _history_text(cid, limit=40)
    if not history.strip():
        raise HTTPException(400, "Belum ada diskusi untuk diringkas")
    fields = await _notulen_fields(u)
    sys = (f"You are the meeting Moderator. You MUST write entirely in {_lang_name(u)}. Write the minutes (notulen) in markdown using EXACTLY these "
           f"sections as headings, in order: {', '.join(f['name'] for f in fields)}. If a section has no content, write '- (belum dibahas)'. "
           "Be concise, warm and neutral.")
    summary = await llm_text(sys, f"Diskusi rapat:\n{history}\n\nRingkasan moderator:")
    used = text_credits(history, summary)
    await record_usage(u["id"], "meeting_summary", used, {"conversation_id": cid})
    msg = {"id": new_id(), "conversation_id": cid, "role": "assistant", "content": summary,
           "persona_id": "__moderator__", "persona_name": "Moderator", "is_moderator": True,
           "portrait": None, "credits": used, "created_at": now_iso()}
    await db.messages.insert_one(dict(msg))
    await db.tasks.insert_one({
        "id": new_id(), "user_id": u["id"], "workspace_id": workspace_id(u), "version": 1, "goal": f"Notulen rapat: {conv['title']}",
        "type": "meeting_notes", "status": "completed", "steps": [], "summary": "Ringkasan & action items panggilan",
        "model": None, "final_output": summary, "credits_used": used, "video_url": None,
        "created_at": now_iso(), "updated_at": now_iso(),
    })
    await db.conversations.update_one({"id": cid}, {"$set": {"updated_at": now_iso()}})
    return {"summary": summary, "message_id": msg["id"], "credits_used": used}


@router.post("/conversations/{cid}/messages/{mid}/regenerate")
async def regenerate(cid: str, mid: str, u: dict = Depends(current_user)):
    conv = await db.conversations.find_one({"id": cid, "user_id": u["id"]})
    if not conv:
        raise HTTPException(404, "Conversation not found")
    old = await db.messages.find_one({"id": mid, "conversation_id": cid}, {"_id": 0})
    if not old:
        raise HTTPException(404, "Message not found")
    persona = await db.personas.find_one({"id": old.get("persona_id")}, {"_id": 0})
    if not persona:
        persona = (await _get_personas(conv.get("persona_ids", [])))[0]
    await db.messages.delete_one({"id": mid, "conversation_id": cid})
    system = await _persona_system(persona, u)
    history = await _history_text(cid)
    full = await llm_text(system, f"{history}\n{persona['name']}:", persona.get("model"))
    used = text_credits(history, full)
    bal = await record_usage(u["id"], "chat", used, {"conversation_id": cid, "regenerate": True})
    ai_msg = {"id": new_id(), "conversation_id": cid, "role": "assistant", "content": full,
              "persona_id": persona["id"], "persona_name": persona["name"],
              "portrait": persona.get("portrait"), "credits": used, "created_at": now_iso()}
    await db.messages.insert_one(dict(ai_msg))
    return {"message": clean(ai_msg), "credits": bal}


# ---------- tools (confirmed by the user) ----------
async def _pending_msg(cid: str, mid: str, u: dict) -> dict:
    conv = await db.conversations.find_one({"id": cid})
    if not _can_access(conv, u):
        raise HTTPException(404, "Conversation not found")
    msg = await db.messages.find_one({"id": mid, "conversation_id": cid}, {"_id": 0})
    if not msg or not msg.get("pending_tool"):
        raise HTTPException(404, "Tidak ada permintaan alat yang menunggu")
    return msg


@router.post("/conversations/{cid}/messages/{mid}/run-tool")
async def run_tool(cid: str, mid: str, app_url: Optional[str] = None, choice: Optional[str] = None, resolution: Optional[str] = None, real_person: bool = False, with_audio: bool = False,
                   aspect: Optional[str] = None, quality: Optional[str] = None, preset: Optional[str] = None, caption: Optional[str] = None, title: Optional[str] = None,
                   u: dict = Depends(current_user)):
    msg = await _pending_msg(cid, mid, u)
    if msg.get("pending_tool", {}).get("kind") == "social":  # caption/title edited on the preview card win over the planner's draft
        pt_edit = {}
        if caption is not None and caption.strip():
            pt_edit["pending_tool.caption"] = caption.strip()[:3000]
        if title is not None and title.strip():
            pt_edit["pending_tool.title"] = title.strip()[:100]
        if pt_edit:
            await db.messages.update_one({"id": mid}, {"$set": pt_edit})
            msg = await _pending_msg(cid, mid, u)
    over = await quota_exceeded(u)
    if over:
        raise HTTPException(402, quota_message(over))
    await rate_limit(u, "generation")
    pt = msg["pending_tool"]
    await db.messages.update_one({"id": mid}, {"$set": {"pending_tool.running": True}})
    if pt.get("kind") == "social":
        try:
            if pt.get("schedule_at"):
                out = await social_schedule_from_chat(cid, u, pt, app_url or "", msg.get("persona_id"), msg.get("persona_name"), msg.get("portrait"))
                upd = {"content": out["text"], "tool": "social_schedule", "social_schedule_id": out["job"]["id"]}
            else:
                out = await social_publish_from_chat(cid, u, pt["providers"], pt.get("caption") or "", pt.get("content_kind") or "image", app_url or "", pt.get("title"))
                upd = {"content": out["text"], "tool": "social_publish"}
        except HTTPException as e:
            upd = {"content": f"Gagal memposting: {e.detail}", "tool": "social_publish", "error": True}
        await db.messages.update_one({"id": mid}, {"$set": upd, "$unset": {"pending_tool": ""}})
        await notify(cid, {"type": "message", "role": "assistant", "persona_id": msg.get("persona_id"), "message": clean({**msg, **upd, "pending_tool": None})})
        return {**msg, **upd, "pending_tool": None}
    if pt.get("kind") == "video":
        duration = int(pt["duration"])
        res = resolution if resolution in seedance.RESOLUTIONS else (pt.get("resolution") or "720p")
        vaspect = aspect if aspect in seedance.ASPECTS else (pt.get("aspect_ratio") or "16:9")
        pt = {**pt, "aspect_ratio": vaspect}
        rp = bool(real_person) and bool(pt.get("reference_path"))
        au = bool(with_audio)
        opt = next((o for o in pt.get("options") or [] if o["tier"] == choice and seedance.supports(o["tier"], duration, res, rp, au)), None)
        if not opt:
            await db.messages.update_one({"id": mid}, {"$unset": {"pending_tool.running": ""}})
            raise HTTPException(400, "Kombinasi model/resolusi/mode ini tidak tersedia — pilih yang lain")
        await pricing_refresh()
        need = seedance.quote(opt["tier"], duration, res, rp, au)
        if await _owner_balance(u) < need:
            await db.messages.update_one({"id": mid}, {"$unset": {"pending_tool.running": ""}})
            raise HTTPException(402, f"Saldo kredit tidak cukup: butuh ±{need} kredit untuk {opt['label']} {duration} detik")
        from integrations import drive_connected
        if not await drive_connected(u["id"]):
            await db.messages.update_one({"id": mid}, {"$unset": {"pending_tool.running": ""}})
            raise HTTPException(400, "Hubungkan Google Drive dulu — video disimpan ke Drive kamu")
        upd = {"content": VIDEO_WAIT_TEXT.replace("render videonya", f"render videonya dengan {opt['label']} {res}{' mode real person' if rp else ''}{' + suara' if au else ''}"), "tool": "video", "rendering": "video", "aspect_ratio": vaspect}
        await db.messages.update_one({"id": mid}, {"$set": upd, "$unset": {"pending_tool": ""}})
        await notify(cid, {"type": "message", "role": "assistant", "persona_id": msg.get("persona_id"), "message": clean({**msg, **upd, "pending_tool": None})})
        asyncio.create_task(_run_video_bg(mid, cid, u["id"], msg.get("persona_id"), {**pt, "tier": opt["tier"], "resolution": res, "real_person": rp, "with_audio": au, "credits": need}, (app_url or "").rstrip("/")))
        return {**msg, **upd, "pending_tool": None}
    if pt.get("kind") == "image_set" and len(pt.get("prompts") or []) > 1:
        persona = await db.personas.find_one({"id": msg.get("persona_id")}, {"_id": 0, "id": 1, "name": 1}) or {"id": msg.get("persona_id"), "name": msg.get("persona_name") or "Asisten"}
        task = await _start_image_set(u, persona, cid, pt["prompts"], f"{len(pt['prompts'])} gambar")
        upd = {"content": _image_set_text(task, len(pt["prompts"])), "tool": "image_set", "task_id": task["id"]}
        await db.messages.update_one({"id": mid}, {"$set": upd, "$unset": {"pending_tool": ""}})
        await notify(cid, {"type": "message", "role": "assistant", "persona_id": msg.get("persona_id"), "message": clean({**msg, **upd, "pending_tool": None})})
        return {**msg, **upd, "pending_tool": None}
    model = choice if choice in {o["id"] for o in pt.get("options") or [] if o.get("available")} else "gemini-image"
    preset = preset if preset in IMAGE_PRESETS else None
    aspect, quality = norm_image_opts(aspect or (IMAGE_PRESETS[preset]["aspect"] if preset else pt.get("aspect_ratio")), quality or pt.get("quality"))
    try:
        out = await run_image_tool(u["id"], pt["prompt"], pt.get("reference_path"), model, aspect, quality, preset)
    except Exception as exc:
        logging.getLogger("chat").warning("run-tool image failed: %s", exc)
        await db.messages.update_one({"id": mid}, {"$unset": {"pending_tool.running": ""}})
        raise HTTPException(502, "Gambar belum berhasil dibuat, coba lagi") from exc
    await record_usage(u["id"], "image_generation" if model == "gemini-image" else "tool:openai:image_generation", out["credits"], {"conversation_id": cid, "persona_id": msg.get("persona_id"), "model": model})
    edited = bool(pt.get("reference_path"))
    out["media"][0]["request"] = pt.get("request") or ""
    upd = {"content": IMAGE_EDIT_DONE_TEXT if edited else IMAGE_DONE_TEXT, "media": out["media"], "tool": "image_edit" if edited else "image", "credits": out["credits"]}
    await db.messages.update_one({"id": mid}, {"$set": upd, "$unset": {"pending_tool": ""}})
    await notify(cid, {"type": "message", "role": "assistant", "persona_id": msg.get("persona_id"), "message": clean({**msg, **upd, "pending_tool": None})})
    return {**msg, **upd, "pending_tool": None}


@router.post("/conversations/{cid}/messages/{mid}/cancel-tool")
async def cancel_tool(cid: str, mid: str, u: dict = Depends(current_user)):
    msg = await _pending_msg(cid, mid, u)
    what = {"video": "video", "social": "posting"}.get((msg.get("pending_tool") or {}).get("kind"), "gambar")
    upd = {"content": f"Oke, pembuatan {what} dibatalkan. Kalau berubah pikiran, tinggal bilang ya!", "tool_cancelled": True}
    await db.messages.update_one({"id": mid}, {"$set": upd, "$unset": {"pending_tool": ""}})
    await notify(cid, {"type": "message", "role": "assistant", "persona_id": msg.get("persona_id"), "message": clean({**msg, **upd, "pending_tool": None})})
    return {**msg, **upd, "pending_tool": None}


# ---------- memory ----------
@router.get("/memory")
async def list_memory(persona_id: Optional[str] = None, u: dict = Depends(current_user)):
    query = {"user_id": u["id"]}
    if persona_id:
        query["persona_id"] = persona_id
    return await db.memory_items.find(query, {"_id": 0}).sort("created_at", -1).to_list(200)


@router.post("/memory")
async def add_memory(x: MemIn, u: dict = Depends(current_user)):
    doc = {"id": new_id(), "user_id": u["id"], "persona_id": x.persona_id,
           "content": x.content, "enabled": True, "created_at": now_iso()}
    await db.memory_items.insert_one(dict(doc))
    return clean(doc)


@router.put("/memory/{mid}")
async def update_memory(mid: str, body: dict, u: dict = Depends(current_user)):
    fields = {}
    if "content" in body:
        fields["content"] = body["content"]
    if "enabled" in body:
        fields["enabled"] = bool(body["enabled"])
    if "pinned" in body:
        fields["pinned"] = bool(body["pinned"])
    await db.memory_items.update_one({"id": mid, "user_id": u["id"]}, {"$set": fields})
    return await db.memory_items.find_one({"id": mid}, {"_id": 0})


@router.delete("/memory/{mid}")
async def del_memory(mid: str, u: dict = Depends(current_user)):
    await db.memory_items.delete_one({"id": mid, "user_id": u["id"]})
    return {"ok": True}
