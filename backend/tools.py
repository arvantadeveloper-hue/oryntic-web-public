import io
import os
import re
import math
import base64
import asyncio
from typing import Optional
from fastapi import APIRouter, Depends
from pydantic import BaseModel, Field

from db import db, now_iso, new_id
from auth import require_platform_admin
from llm import llm_text, llm_json, generate_image, text_credits, _MODEL_BY_ID, model_label
from pricing import rate
from storage import put_object

router = APIRouter(prefix="/api", tags=["tools"])

DEFAULT_ROUTING = {"enabled": True, "it_model": "claude-sonnet", "research_model": "gemini-pro", "confirm_threshold": 20}
_routing_cache = {"doc": None, "at": 0.0}

IT_RE = re.compile(r"\b(kode|code|coding|bug|error|exception|stack ?trace|python|javascript|typescript|react|node|sql|query|api|endpoint|server|deploy|docker|kubernetes|database|regex|git|html|css|fungsi|function|script|debug|compile|library|framework|algoritma|algorithm|json|yaml|linux|terminal|command|npm|pip)\b", re.I)
RESEARCH_RE = re.compile(r"\b(riset|research|penelitian|analisis mendalam|analisa mendalam|literatur|kajian|bandingkan secara|studi kasus|laporan lengkap|ringkas(kan)? dokumen|whitepaper|jurnal|tinjauan)\b", re.I)
TOOL_RE = re.compile(r"\b(gambar|image|ilustrasi|foto|logo|poster|banner|visual|sketsa|lukisan|dokumen|document|laporan|proposal|docx|pdf|word|file|surat|artikel|template|makalah|ringkasan tertulis|notulen)\b", re.I)
CREATE_RE = re.compile(r"\b(buat|buatkan|bikin|bikinkan|generate|tolong buat|susun|susunkan|tuliskan|rancang|desain|design|create|make|draw|gambarkan|ekspor|export|unduh|download)\b", re.I)


async def get_routing() -> dict:
    import time
    if _routing_cache["doc"] and time.time() - _routing_cache["at"] < 30:
        return _routing_cache["doc"]
    cfg = await db.config.find_one({"id": "model_routing"}, {"_id": 0, "id": 0, "updated_at": 0}) or {}
    doc = {**DEFAULT_ROUTING, **cfg}
    _routing_cache.update(doc=doc, at=time.time())
    return doc


class RoutingIn(BaseModel):
    enabled: bool = True
    it_model: str = Field(pattern="^[a-z0-9-]+$")
    research_model: str = Field(pattern="^[a-z0-9-]+$")
    confirm_threshold: int = Field(ge=0, le=1000)


@router.get("/admin/model-routing")
async def admin_get_routing(_: dict = Depends(require_platform_admin)):
    return {**(await get_routing()), "models": [{"id": m["id"], "label": m["label"]} for m in _MODEL_BY_ID.values()]}


@router.put("/admin/model-routing")
async def admin_set_routing(x: RoutingIn, _: dict = Depends(require_platform_admin)):
    doc = x.model_dump()
    if doc["it_model"] not in _MODEL_BY_ID or doc["research_model"] not in _MODEL_BY_ID:
        from fastapi import HTTPException
        raise HTTPException(400, "Model tidak dikenal")
    await db.config.update_one({"id": "model_routing"}, {"$set": {**doc, "updated_at": now_iso()}}, upsert=True)
    _routing_cache["doc"] = None
    return await admin_get_routing(_)


async def route_model(persona_model: Optional[str], text: str, extra_len: int, owner_settings: dict):
    """Pick the best model for this turn while the persona stays the same. Returns (model_key, reason|None)."""
    cfg = await get_routing()
    if not cfg["enabled"] or not (owner_settings or {}).get("smart_routing", True):
        return persona_model, None
    t = text or ""
    if IT_RE.search(t) or "```" in t:
        return cfg["it_model"], "it"
    if RESEARCH_RE.search(t) or len(t) + extra_len > 3500:
        return cfg["research_model"], "research"
    return persona_model, None


DRIVE_RE = re.compile(r"\b(google ?drive|drive|gdrive|google ?docs?|spreadsheet|google ?sheets?)\b", re.I)
SOCIAL_RE = re.compile(r"\b(linkedin|instagram|facebook|youtube|sosmed|social media|medsos|posting|post(kan)?|unggah ke)\b", re.I)
GITHUB_RE = re.compile(r"\b(github|gitlab|repo|repository|repositori|pull ?request|merge ?request|PR|MR|issues?|branch|commit|review|diff)\b", re.I)


MEDIA_RE = re.compile(r"\b(gambar\w*|image|picture|foto\w*|photo\w*|ilustrasi|illustration|logo|poster|banner|visual\w*|sketsa|lukisan|render\w*|wallpaper|thumbnail|video\w*|klip|clip|animasi|animation|reels?|footage|cuplikan)\b", re.I)
REFUSAL_RE = re.compile(r"(belum|tidak|nggak|gak|tak)\s+(bisa|dapat|mampu|sanggup)\s+(me(m|n|ng)?)?(render|buat|bikin|hasilkan|tampilkan|munculkan|generate|kirim)\w*[^.\n]{0,40}\b(gambar|foto|image|visual|video|klip|ilustrasi)|\b(can(no|')t|unable to)\s+(render|generate|create|display|make|show)\w*[^.\n]{0,40}\b(image|picture|photo|video|visual)", re.I)


def wants_tool(text: str) -> bool:
    return bool(text) and ((bool(TOOL_RE.search(text)) and bool(CREATE_RE.search(text))) or bool(MEDIA_RE.search(text)) or bool(DRIVE_RE.search(text)) or bool(GITHUB_RE.search(text)) or bool(SOCIAL_RE.search(text)))


def is_media_refusal(user_text: str, reply: str) -> bool:
    """The model claimed it cannot render an image/video although the user asked for one."""
    return bool(user_text and reply and MEDIA_RE.search(user_text) and REFUSAL_RE.search(reply))


async def plan_tool(text: str, history: str) -> dict:
    sys = ('Decide if the user\'s LAST message explicitly asks the assistant to CREATE a deliverable or act on an external service. Reply JSON only: '
           '{"tool":"image"|"image_edit"|"video"|"document"|"drive_save"|"drive_update"|"drive_link"|"github_repos"|"github_read"|"github_issues"|"github_pr"|"github_review"|"gitlab_repos"|"gitlab_read"|"gitlab_issues"|"gitlab_pr"|"gitlab_review"|"social_publish"|"none",'
           '"image_prompt":str,"image_prompts":[str],"quality":"hemat"|"standar"|"tinggi","video_prompt":str,"duration":int,"aspect_ratio":"16:9"|"9:16"|"1:1"|"4:3"|"3:4"|"21:9","resolution":"480p"|"720p"|"1080p","real_person":bool,"with_audio":bool,"from_image":bool,"edit_prompt":str,"title":str,"instructions":str,"file":str,"text":str,"mode":"append"|"replace","kind":"doc"|"sheet",'
           '"repo":str,"path":str,"query":str,"state":"open"|"closed"|"all","files":[str],"number":int,"providers":[str],"caption":str,"content_kind":"text"|"image"|"video"}. '
           '"image" = the user wants ANY still visual generated, shown or rendered: photo, photorealistic/realistic picture, render, illustration, logo, poster, banner, wallpaper, thumbnail, sketch, painting, visualization ("tunjukkan", "tampilkan", "render", "visualisasikan", "gambarkan" count as a request). If they ask for MORE THAN ONE image (e.g. "3 variasi", "beberapa poster", '
           '"gambar A dan gambar B"), put one detailed English prompt PER image in image_prompts (max 6) and the first one in image_prompt; for a single image image_prompts has exactly one item. '
           'For image/image_edit also set aspect_ratio: "9:16" for portrait/vertical/story/poster/phone wallpaper ("potret", "vertikal", "tegak"), "16:9" for landscape/wide/banner/desktop wallpaper/thumbnail/presentation ("lanskap", "melebar", "horizontal"), "4:3"/"3:4" only when explicitly asked, else "1:1"; '
           'and quality: "hemat" when they ask for cheap/quick/draft ("hemat", "murah", "cepat", "draf"), "tinggi" for high quality/HD/detailed/print/4K ("kualitas tinggi", "tajam", "detail", "HD", "4K", "cetak"), else "standar". '
           '"image_edit" = the user asks to CHANGE the image the assistant generated earlier in this conversation — restyle ("ubah jadi gaya kartun/anime/lukisan cat air"), '
           'recolor, change background/lighting/time of day, add or remove an object, make it brighter/darker, crop, make a variation that keeps the same subject. '
           'edit_prompt: a precise English editing instruction for an image model that receives the previous image as reference (describe what to change and what must stay the same). '
           'Use "image" (not image_edit) when they clearly want a brand-new picture of something else. '
           '"video" = the user wants a short video/clip/animation/reel/footage generated or rendered (video_prompt: detailed English prompt describing scene, motion, camera, mood; duration: whole seconds the user asked for, 4–30, else 5; aspect_ratio: "9:16" when they say portrait/vertical/Reels/Shorts/TikTok/story, "1:1" for square/feed Instagram, "21:9" for ultra-wide cinematic, "4:3"/"3:4" only if explicitly asked, else "16:9"; resolution: "480p" when they ask for cheap/economical/low-res ("hemat", "murah", "ringan"), "1080p" for sharp/HD/full HD/high quality ("tajam", "HD", "1080"), else "720p"; real_person: true when they want real-person mode / a real human face preserved ("mode real person", "orang asli", "wajah asli") — only meaningful together with from_image; with_audio: true when they want sound/ambience/sound effects/background audio generated with the clip ("dengan suara", "ada suaranya", "pakai ambience", "with sound"); from_image: true when they want the image the assistant generated earlier to be animated / turned into a video — "animasikan gambar ini", "jadikan video", "gerakkan gambarnya", "buat videonya dari gambar tadi"). '
           '"document" = the user wants a written file '
           '(report, proposal, letter, article, notulen, template) they can download. Otherwise "none" (questions, explanations, '
           'tables shown inline, code snippets are NOT documents). image_prompt: detailed English prompt for an image model. '
           'title: short document title in the user\'s language. instructions: what the document must contain. '
           '"drive_save" = user asks to save/store something (the current document, a table, the discussed content) to Google Drive/Docs/Sheets '
           '(kind "sheet" when they want a spreadsheet, else "doc"; text = the content to save if they described it). '
           '"drive_update" = user asks to update/edit/append to an existing Google Doc on Drive (file = document name or id as the user said it; '
           'text = the exact content to add/write; mode "replace" only if they want to overwrite). '
           '"drive_link" = user asks for the Drive link of a document (file = name). '
           '"github_repos" = user asks to list/find their GitHub repositories (query = optional filter). '
           '"github_read" = user asks to read/show/explain a file, folder or the structure of a GitHub repo (repo = owner/name as mentioned or from context; path = file/folder path or "" for the tree). '
           '"github_issues" = user asks about issues/PRs of a repo (repo; state). '
           '"github_pr" = user asks to change code/files in a repo and open a pull request / make a PR / fix something in the repo (repo; instructions = what to change; files = file paths they mentioned, may be empty; title = short PR title). '
           '"github_review" = user asks to review/check/evaluate/summarize a pull request or MR, its changes or diff (repo; number = PR/MR number if mentioned, e.g. "#12" or "!12", else omit → latest open one). '
           'Use gitlab_* (same meanings; gitlab_pr = open a Merge Request, gitlab_review = review an MR) when the user says GitLab / merge request / MR or the project is known to be on GitLab; otherwise github_*. '
           'Only use github_*/gitlab_* when a code repository, PR/MR or issue is clearly meant. '
           '"social_publish" = user asks to post/publish/share content to social media (providers from: linkedin, meta (Facebook Page/Instagram), youtube — map Instagram/Facebook→meta; content_kind: image = the latest generated image, video = the latest video, text = a text-only post; caption = the caption they gave, or write a fitting one in their language).')
    plan: dict = {}
    try:
        plan = await llm_json(sys, f"Recent conversation:\n{history[-2500:]}\n\nLAST MESSAGE: {text}")
    except Exception:
        plan = {}
    if plan.get("tool") not in ("image", "image_edit", "video", "document", "drive_save", "drive_update", "drive_link", "github_repos", "github_read", "github_issues", "github_pr", "github_review", "gitlab_repos", "gitlab_read", "gitlab_issues", "gitlab_pr", "gitlab_review", "social_publish"):
        return {"tool": "none"}
    if plan["tool"] == "image":
        prompts = [str(p).strip() for p in (plan.get("image_prompts") or []) if str(p).strip()][:6]
        plan["image_prompts"] = prompts or [(plan.get("image_prompt") or "").strip() or "illustration"]
        plan["image_prompt"] = plan["image_prompts"][0]
    if plan["tool"] == "image_edit":
        plan["edit_prompt"] = (plan.get("edit_prompt") or plan.get("image_prompt") or text).strip()
    if plan["tool"] in ("image", "image_edit"):
        plan["aspect_ratio"], plan["quality"] = norm_image_opts(plan.get("aspect_ratio"), plan.get("quality"))
    if plan["tool"] == "video":
        plan["video_prompt"] = (plan.get("video_prompt") or plan.get("image_prompt") or "").strip() or "cinematic short clip"
        from seedance import clamp_duration
        plan["duration"] = clamp_duration(plan.get("duration"))
        plan["from_image"] = bool(plan.get("from_image"))
        from seedance import ASPECTS, RESOLUTIONS
        plan["aspect_ratio"] = plan.get("aspect_ratio") if plan.get("aspect_ratio") in ASPECTS else "16:9"
        plan["resolution"] = plan.get("resolution") if plan.get("resolution") in RESOLUTIONS else "720p"
        plan["real_person"] = bool(plan.get("real_person")) and plan["from_image"]
        plan["with_audio"] = bool(plan.get("with_audio"))
    return plan


def _md_blocks(md: str):
    for raw in md.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            yield ("blank", "")
        elif line.startswith("```"):
            yield ("fence", "")
        elif re.match(r"^#{1,3}\s", line):
            yield ("h%d" % len(re.match(r"^#+", line).group()), re.sub(r"^#+\s", "", line))
        elif re.match(r"^\s*[-*]\s", line):
            yield ("li", re.sub(r"^\s*[-*]\s", "", line))
        elif re.match(r"^\s*\d+\.\s", line):
            yield ("li", re.sub(r"^\s*", "", line))
        elif re.match(r"^\s*\|.*\|\s*$", line):
            if not re.match(r"^\s*\|[\s:|-]+\|\s*$", line):
                yield ("row", [c.strip() for c in line.strip()[1:-1].split("|")])
        else:
            yield ("p", line)


_inline = re.compile(r"\*\*([^*]+)\*\*|\*([^*]+)\*|`([^`]+)`")


def _plain(s: str) -> str:
    return _inline.sub(lambda m: m.group(1) or m.group(2) or m.group(3), s)


def build_docx(title: str, md: str) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading(title, 0)
    table_rows = []

    def flush_table():
        if not table_rows:
            return
        t = doc.add_table(rows=0, cols=max(len(r) for r in table_rows))
        t.style = "Table Grid"
        for r in table_rows:
            cells = t.add_row().cells
            for i, c in enumerate(r):
                cells[i].text = _plain(c)
        table_rows.clear()

    for kind, val in _md_blocks(md):
        if kind != "row":
            flush_table()
        if kind == "row":
            table_rows.append(val)
        elif kind.startswith("h"):
            doc.add_heading(_plain(val), int(kind[1]))
        elif kind == "li":
            doc.add_paragraph(_plain(val), style="List Bullet")
        elif kind == "p":
            doc.add_paragraph(_plain(val))
    flush_table()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pdf(title: str, md: str) -> bytes:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    reg, bold = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
    fam = "Helvetica"
    if os.path.exists(reg):
        pdf.add_font("Lib", "", reg)
        pdf.add_font("Lib", "B", bold if os.path.exists(bold) else reg)
        fam = "Lib"
    w = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.set_font(fam, "B", 18)
    pdf.multi_cell(w, 9, title)
    pdf.ln(2)
    rows = []

    def flush_rows():
        if not rows:
            return
        cols = max(len(r) for r in rows)
        cw = w / cols
        for ri, r in enumerate(rows):
            pdf.set_font(fam, "B" if ri == 0 else "", 9)
            cells = [_plain(r[i]) if i < len(r) else "" for i in range(cols)]
            h = 5 * max(1, max(len(pdf.multi_cell(cw - 2, 5, c, dry_run=True, output="LINES")) for c in cells))
            if pdf.get_y() + h > pdf.page_break_trigger:
                pdf.add_page()
            y0 = pdf.get_y()
            for i, c in enumerate(cells):
                x = pdf.l_margin + i * cw
                pdf.rect(x, y0, cw, h)
                pdf.set_xy(x + 1, y0)
                pdf.multi_cell(cw - 2, 5, c, border=0, new_x="RIGHT", new_y="TOP")
            pdf.set_xy(pdf.l_margin, y0 + h)
        rows.clear()
        pdf.ln(2)

    for kind, val in _md_blocks(md):
        if kind != "row":
            flush_rows()
        if kind == "row":
            rows.append(val)
        elif kind.startswith("h"):
            pdf.ln(2)
            pdf.set_font(fam, "B", {1: 15, 2: 13, 3: 11}[int(kind[1])])
            pdf.multi_cell(w, 7, _plain(val))
        elif kind == "li":
            pdf.set_font(fam, "", 10.5)
            pdf.multi_cell(w, 6, "•  " + _plain(val))
        elif kind == "p":
            pdf.set_font(fam, "", 10.5)
            pdf.multi_cell(w, 6, _plain(val))
        elif kind == "blank":
            pdf.ln(2)
    flush_rows()
    return bytes(pdf.output())


def _safe_name(title: str) -> str:
    s = re.sub(r"[^\w\- ]+", "", title or "dokumen").strip().replace(" ", "_")
    return (s or "dokumen")[:60]


EDIT_RE = re.compile(r"\b(ubah|ganti|rubah|jadikan|bikin jadi|buat jadi|tambah(kan|in)?|hapus|hilangkan|kurangi|perbesar|perkecil|lebih (terang|gelap|cerah|tajam|halus)|versi|gaya|style|warna|latar|background|edit|variasi|crop|potong|zoom|make it|change|turn it|add|remove|more|less)\b", re.I)


IMAGE_MODELS = {
    "gemini-image": {"label": "Nano Banana (Gemini)", "desc": "Cepat, cocok untuk ilustrasi & edit foto"},
    "gpt-image": {"label": "GPT Image (OpenAI)", "desc": "Detail tinggi, teks di gambar lebih akurat"},
}


IMAGE_ASPECTS = {"1:1": "Persegi", "16:9": "Lanskap", "9:16": "Potret", "4:3": "Lanskap 4:3", "3:4": "Potret 3:4"}
IMAGE_QUALITIES = ("hemat", "standar", "tinggi")
# quality → (provider setting, price multiplier vs. the model's base price); gpt-image non-square sizes cost ×1.5
_QUALITY = {"gemini-image": {"hemat": ("1K", 1.0), "standar": ("2K", 1.5), "tinggi": ("4K", 2.25)},
            "gpt-image": {"hemat": ("low", 0.27), "standar": ("medium", 1.0), "tinggi": ("high", 4.0)}}
_GPT_SIZE = {"1:1": "1024x1024", "16:9": "1536x1024", "4:3": "1536x1024", "9:16": "1024x1536", "3:4": "1024x1536"}


def norm_image_opts(aspect: Optional[str], quality: Optional[str]) -> tuple:
    return (aspect if aspect in IMAGE_ASPECTS else "1:1", quality if quality in IMAGE_QUALITIES else "standar")


def image_quote(model: str, base: int, aspect: str, quality: str) -> int:
    mult = _QUALITY[model][quality][1] * (1.5 if model == "gpt-image" and aspect != "1:1" else 1.0)
    return max(1, math.ceil(base * mult))


async def _image_base_credits() -> dict:
    from pricing import get_pricing, tool_credits
    return {"gemini-image": rate("image"), "gpt-image": tool_credits(await get_pricing(), "openai:image_generation")}


async def image_model_options() -> list:
    """Models offered on the image confirmation card: base credits, availability and the credits matrix per quality × aspect."""
    base = await _image_base_credits()
    gpt_ok = bool(os.environ.get("OPENAI_API_KEY", "").strip())
    return [{"id": m, **IMAGE_MODELS[m], "credits": base[m], "available": m != "gpt-image" or gpt_ok,
             "prices": {q: {a: image_quote(m, base[m], a, q) for a in IMAGE_ASPECTS} for q in IMAGE_QUALITIES}} for m in IMAGE_MODELS]


async def _gpt_image(prompt: str, ref_b64: Optional[str], size: str, quality: str) -> str:
    """OpenAI gpt-image-1 (BYOK) → data URL. With a reference → edit, else generate."""
    from openai import AsyncOpenAI
    client = AsyncOpenAI(api_key=os.environ["OPENAI_API_KEY"])
    if ref_b64:
        import io
        f = io.BytesIO(base64.b64decode(ref_b64)); f.name = "reference.png"
        r = await client.images.edit(model="gpt-image-1", image=f, prompt=prompt[:4000], size=size, quality=quality)
    else:
        r = await client.images.generate(model="gpt-image-1", prompt=prompt[:4000], size=size, quality=quality)
    return "data:image/png;base64," + r.data[0].b64_json


async def run_image_tool(uid: str, prompt: str, reference_path: Optional[str] = None, model: str = "gemini-image", aspect: str = "1:1", quality: str = "standar") -> dict:
    ref_b64 = None
    if reference_path:
        from storage import get_object
        raw_ref, _ct = await asyncio.to_thread(get_object, reference_path)
        ref_b64 = base64.b64encode(raw_ref).decode()
    model = model if model in IMAGE_MODELS else "gemini-image"
    aspect, quality = norm_image_opts(aspect, quality)
    setting = _QUALITY[model][quality][0]
    if model == "gpt-image":
        data_url = await _gpt_image(prompt, ref_b64, _GPT_SIZE[aspect], setting)
    else:
        data_url = await generate_image(prompt, ref_b64, aspect_ratio=aspect, image_size=setting)
    credits = image_quote(model, (await _image_base_credits())[model], aspect, quality)
    if not data_url:
        raise RuntimeError("image generation returned nothing")
    header, b64 = data_url.split(",", 1)
    mime = header.split(":")[1].split(";")[0]
    ext = "png" if "png" in mime else "jpg"
    path = f"aivora/images/{uid}/{new_id()}.{ext}"
    from integrations import assert_quota, add_storage
    raw = base64.b64decode(b64)
    await assert_quota(uid, len(raw))
    await asyncio.to_thread(put_object, path, raw, mime)
    await add_storage(uid, len(raw))
    media = {"type": "image", "path": path, "name": f"gambar.{ext}", "prompt": prompt[:400], "model": IMAGE_MODELS[model]["label"], "aspect_ratio": aspect, "quality": quality}
    if reference_path:
        media["edited_from"] = reference_path
    return {"media": [media], "credits": credits}


async def run_document_tool(uid: str, system: str, title: str, instructions: str, history: str, model_key: Optional[str]) -> dict:
    sys = system + ("\n\nDOCUMENT MODE: write the COMPLETE document content in well-structured markdown (headings, lists, tables "
                    "where useful). No preface, no closing chit-chat — only the document body. The whole document follows the LANGUAGE RULE above "
                    "(the user's language) unless the user explicitly asked for the document in another language.")
    prompt = f"Context:\n{history[-3000:]}\n\nDocument title: {title}\nRequirements: {instructions}\n\nWrite the full document now, in the user's language per the LANGUAGE RULE."
    md = await llm_text(sys, prompt, model_key)
    base = _safe_name(title)
    docx_b, pdf_b = await asyncio.gather(asyncio.to_thread(build_docx, title, md), asyncio.to_thread(build_pdf, title, md))
    did = new_id()
    files = [("docx", docx_b, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
             ("pdf", pdf_b, "application/pdf"), ("md", md.encode("utf-8"), "text/markdown")]
    media = []
    from integrations import assert_quota, add_storage
    await assert_quota(uid, sum(len(d) for _, d, _ in files))
    for ext, data, ctype in files:
        path = f"aivora/docs/{uid}/{did}/{base}.{ext}"
        await asyncio.to_thread(put_object, path, data, ctype)
        await add_storage(uid, len(data))
        media.append({"type": "file", "path": path, "name": f"{base}.{ext}", "format": ext})
    return {"media": media, "credits": text_credits(prompt, md, model_key), "markdown": md, "model_label": model_label(model_key)}


# ---------- workspace tasks: shared context & revisions ----------
TABLE_RE = re.compile(r"^\s*\|.+\|\s*$", re.M)
TASK_CONTEXT = ("\n\nWORKSPACE TASK UNDER DISCUSSION (id {tid}, version {ver}, status {status}; the ONLY valid link to it is /workspace/{tid} — never invent other URLs):\nTitle: {goal}\n--- CURRENT RESULT ---\n{body}\n--- END ---\n"
                "The user may ask questions about this result or request changes. Discuss it in context; when they ask for a change, the system "
                "saves a revised version to the Workspace automatically — confirm briefly what changed.")
REVISE_RE = re.compile(r"\b(revisi|ubah|ganti|perbaiki|tambah(kan)?|hapus|kurangi|perbarui|update|rapikan|singkat|perpanjang|sesuaikan|koreksi|edit|rewrite|revise|change|tulis ulang|terapkan|simpan|masukkan|pakai|gunakan|jadikan|replace|apply|save)\b", re.I)


def has_tables(md: str) -> bool:
    return len(TABLE_RE.findall(md or "")) >= 2



async def save_revision(task: dict, new_md: str, note: str, persona: Optional[dict]) -> int:
    """Archive the current result as a version and store the revised text. Returns the new version number."""
    cur = int(task.get("version") or 1)
    old = {"version": cur, "content": task.get("final_output") or "", "created_at": task.get("updated_at") or task.get("created_at"), "note": task.get("revision_note") or "Versi awal"}
    await db.tasks.update_one({"id": task["id"]}, {"$push": {"versions": old},
                                                   "$set": {"final_output": new_md, "version": cur + 1, "revision_note": note[:300], "updated_at": now_iso(),
                                                            "status": "completed", **({"persona_id": persona["id"], "persona_name": persona["name"]} if persona else {})}})
    return cur + 1


async def revise_with_llm(task: dict, request: str, system: str, model_key: Optional[str], history: str = "") -> tuple:
    """Returns (new_markdown, change_summary, credits). `history` = recent chat so "pakai konten di atas" resolves to the right text."""
    sys = system + ("\n\nYou are REVISING a workspace deliverable. Output the COMPLETE revised document in markdown (keep everything that was "
                    "not asked to change), then on the very last line write exactly: RINGKASAN PERUBAHAN: <1-2 sentences>. Keep the document in the user's language "
                    "per the LANGUAGE RULE above — never translate it to another language unless the revision request explicitly asks for that.")
    ctx = f"RECENT CONVERSATION (the request may refer to content proposed here — when the user approves content from this conversation, use it verbatim as the new document body/section):\n{history[-12000:]}\n\n" if history else ""
    prompt = f"{ctx}CURRENT DOCUMENT:\n{(task.get('final_output') or '')[:40000]}\n\nREVISION REQUEST: {request}"
    out = await llm_text(sys, prompt, model_key)
    summary = ""
    if "RINGKASAN PERUBAHAN:" in out:
        out, summary = out.rsplit("RINGKASAN PERUBAHAN:", 1)
    return out.strip(), summary.strip(), text_credits(prompt, out, model_key)




# ---------- task assignment from chat / meeting ----------
TASK_RE = re.compile(r"\b(tolong|bisa|minta|buatkan|buatlah|kerjakan|susun(kan)?|siapkan|rancang|analisis|analisa|riset|teliti|rangkum|ringkas|terjemahkan|"
                     r"laporan|proposal|rencana|roadmap|strategi|artikel|esai|modul|kurikulum|presentasi|jadwalkan|nanti|besok|lusa|minggu depan|jam \d|pukul \d|deadline|tenggat)\b", re.I)
OFFER_TEXT = ("Siap, {uname}! Tugas ini cukup besar: **{title}**.\n\nMau kita **bahas satu per satu** di sini, atau **terima beres** saja? "
              "Kalau terima beres, tugas saya masukkan ke Ruang Kerja dan kerjakan {when}; setelah selesai saya kabari di chat ini, dan Anda bisa buat panggilan supaya saya paparkan hasilnya.")


async def plan_task(text: str, tz: str, history: str = "") -> dict:
    """Decide whether the message delegates a (long) piece of work and when it should be done."""
    if not TASK_RE.search(text or "") or len(text) < 25:
        return {"is_task": False}
    from zoneinfo import ZoneInfo
    from datetime import datetime
    try:
        now = datetime.now(ZoneInfo(tz or "Asia/Jakarta"))
    except Exception:
        now = datetime.now(ZoneInfo("Asia/Jakarta"))
    try:
        r = await llm_json(
            "You classify whether a user's chat message DELEGATES work to the assistant (a deliverable: document, plan, analysis, research, report, "
            "code module, etc.) versus a quick question. Reply JSON only: {\"is_task\": bool, \"long\": bool, \"title\": str, \"brief\": str, \"scheduled_at\": str|null}. "
            "long=true when the deliverable needs more than ~2 paragraphs of real work. title: short Indonesian title (max 10 words). brief: 1-3 sentences "
            "restating exactly what must be produced. scheduled_at: ISO-8601 with timezone offset if the user names a time/day to do it (e.g. 'besok jam 9', "
            "'Senin depan'), else null.",
            f"Now: {now.isoformat()} ({tz}).\nRecent context: {history[-600:]}\nUser message: {text}")
    except Exception:
        return {"is_task": False}
    return r if isinstance(r, dict) else {"is_task": False}


CAL_RE = re.compile(r"\b(catat(kan)?|jadwalkan|ingatkan|tambahkan|masukkan|simpan|buat(kan)? pengingat|set pengingat)\b.*\b(kalender|agenda|pengingat|jadwal|rapat|meeting|janji|acara|kegiatan|deadline|tenggat|ulang tahun|appointment|reminder)\b"
                    r"|\b(ingatkan|remind) (saya|aku|gue|gw|me)\b", re.I | re.S)


async def plan_calendar(text: str, tz: str, history: str = "", force: bool = False) -> dict:
    """Extract a calendar entry (+ reminder preferences) from a chat message. Returns {} when it is not a calendar request."""
    if not force and not CAL_RE.search(text or ""):
        return {}
    from zoneinfo import ZoneInfo
    from datetime import datetime
    try:
        now = datetime.now(ZoneInfo(tz or "Asia/Jakarta"))
    except Exception:
        now = datetime.now(ZoneInfo("Asia/Jakarta"))
    try:
        r = await llm_json(
            "You extract a CALENDAR ENTRY the user wants recorded (meeting, appointment, activity, deadline, birthday, reminder). Reply JSON only: "
            "{\"is_calendar\": bool, \"title\": str, \"start_at\": str|null, \"notes\": str, \"remind_mode\": \"call\"|\"chat\"|null, \"remind_offsets\": [int], \"repeat\": \"none\"|\"daily\"|\"weekly\"|\"monthly\", \"question\": str|null}. "
            "title: short Indonesian title (max 8 words, no date words). start_at: ISO-8601 WITH timezone offset, resolved from relative words ('besok jam 10', 'Jumat depan sore') using Now; "
            "null when no usable time was given. notes: extra details (place, people, agenda) or ''. remind_mode: 'call' if the user wants to be called/phoned, 'chat' if via message/chat, "
            "null if unspecified. remind_offsets: minutes before start the user asked for (e.g. '30 menit dan 1 jam sebelum' → [30, 60]); [] if unspecified; use [0] when the user explicitly wants NO reminder. "
            "repeat: 'daily'/'weekly'/'monthly' when the user says setiap hari/minggu/bulan, tiap Senin, bulanan, rutin…; else 'none'. For repeating entries start_at is the FIRST occurrence (for 'tiap Senin jam 8' → next Monday 08:00). "
            "question: when start_at is missing or truly ambiguous (date without time, or 'minggu depan' without a day), ONE short Indonesian question asking exactly what is missing; else null. "
            "is_calendar=false for anything that is not a request to record/schedule/remind.",
            f"Now: {now.isoformat()} ({tz}).\nRecent context: {history[-600:]}\nUser message: {text}")
    except Exception:
        return {}
    return r if isinstance(r, dict) and r.get("is_calendar") else {}
