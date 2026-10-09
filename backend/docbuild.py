"""Document builders: markdown → DOCX / PDF bytes (used by the document tool in tools.py)."""
import io
import os
import re

from mdblocks import md_blocks


_inline = re.compile(r"\*\*([^*]+)\*\*|\*([^*]+)\*|`([^`]+)`")


def _plain(s: str) -> str:
    return _inline.sub(lambda m: m.group(1) or m.group(2) or m.group(3), s)


def build_docx(title: str, md: str) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading(title, 0)
    table_rows = []

    def flush_table():
        if not table_rows:
            return
        t = doc.add_table(rows=0, cols=max(len(r) for r in table_rows))
        t.style = "Table Grid"
        for r in table_rows:
            cells = t.add_row().cells
            for i, c in enumerate(r):
                cells[i].text = _plain(c)
        table_rows.clear()

    for kind, val in md_blocks(md):
        if kind != "row":
            flush_table()
        if kind == "row":
            table_rows.append(val)
        elif kind.startswith("h"):
            doc.add_heading(_plain(val), int(kind[1]))
        elif kind == "li":
            doc.add_paragraph(_plain(val), style="List Bullet")
        elif kind == "p":
            doc.add_paragraph(_plain(val))
    flush_table()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pdf(title: str, md: str) -> bytes:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    reg, bold = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
    fam = "Helvetica"
    if os.path.exists(reg):
        pdf.add_font("Lib", "", reg)
        pdf.add_font("Lib", "B", bold if os.path.exists(bold) else reg)
        fam = "Lib"
    w = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.set_font(fam, "B", 18)
    pdf.multi_cell(w, 9, title)
    pdf.ln(2)
    rows = []

    def flush_rows():
        if not rows:
            return
        cols = max(len(r) for r in rows)
        cw = w / cols
        for ri, r in enumerate(rows):
            pdf.set_font(fam, "B" if ri == 0 else "", 9)
            cells = [_plain(r[i]) if i < len(r) else "" for i in range(cols)]
            h = 5 * max(1, max(len(pdf.multi_cell(cw - 2, 5, c, dry_run=True, output="LINES")) for c in cells))
            if pdf.get_y() + h > pdf.page_break_trigger:
                pdf.add_page()
            y0 = pdf.get_y()
            for i, c in enumerate(cells):
                x = pdf.l_margin + i * cw
                pdf.rect(x, y0, cw, h)
                pdf.set_xy(x + 1, y0)
                pdf.multi_cell(cw - 2, 5, c, border=0, new_x="RIGHT", new_y="TOP")
            pdf.set_xy(pdf.l_margin, y0 + h)
        rows.clear()
        pdf.ln(2)

    for kind, val in md_blocks(md):
        if kind != "row":
            flush_rows()
        if kind == "row":
            rows.append(val)
        elif kind.startswith("h"):
            pdf.ln(2)
            pdf.set_font(fam, "B", {1: 15, 2: 13, 3: 11}[int(kind[1])])
            pdf.multi_cell(w, 7, _plain(val))
        elif kind == "li":
            pdf.set_font(fam, "", 10.5)
            pdf.multi_cell(w, 6, "•  " + _plain(val))
        elif kind == "p":
            pdf.set_font(fam, "", 10.5)
            pdf.multi_cell(w, 6, _plain(val))
        elif kind == "blank":
            pdf.ln(2)
    flush_rows()
    return bytes(pdf.output())
